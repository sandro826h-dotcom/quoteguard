"""由初标 JSON 生成复核包：每单元一个目录，含匿名原件、初标 JSON 与逐字段复核表（xlsx，黄色列由复核人填写）。

用法：python3 scripts/make_review_pack.py <初标目录> <输出目录>
"""
from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
ANON = ROOT / "data" / "seeds" / "anonymous"
FIELDS = ["channel", "mode", "rate_tiers", "min_kg_shipment", "min_kg_carton", "dim_divisor", "rounding_kg",
          "surcharges", "fuel_pct", "includes", "excludes", "ddp", "broker_fee_cny", "isf_fee_cny",
          "transit_days", "validity", "liability", "claims"]
YELLOW = PatternFill("solid", fgColor="FFF2CC")
HEAD = ["字段", "status（初标）", "value（初标）", "出处片段（初标）", "初标备注",
        "确认/改正", "改正后 status", "改正值", "原文依据", "理由"]
WIDTH = [16, 12, 48, 60, 40, 12, 14, 36, 40, 30]


def compact(v) -> str:
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    return json.dumps(v, ensure_ascii=False)


def spans_text(spans) -> str:
    return "\n".join(f"[{s['doc_id']}{' p' + str(s['page']) if s.get('page') else ''}] {s['quote']}" for s in spans)


def field_note(notes: dict, name: str) -> str:
    keys = [k for k in notes if k == name or k.startswith(name + "_")]
    if name == "rate_tiers":
        keys += [k for k in ("range_max", "price_form") if k in notes]
    return "\n".join(f"{k}: {compact(notes[k])}" for k in keys)


def sheet(wb: Workbook, rec: dict) -> None:
    ws = wb.active
    ws.title = "逐字段复核"
    m = rec["meta"]
    ws.append([f"{rec['seed_id']}  渠道：{m['channel']}  区域：{m['dest_region']}  原件：{m['source_id']}"
               f"（{m['source_type']}）  price_form：{m['price_form']}  date_status：{m['date_status']}"])
    ws.append([f"初标：{rec['initial_label_model']}。黄色列由复核人填写：确认/改正 选“确认”或“改正”；改正时填改正后 status、"
               "改正值（按规范第六节归一化）、原文依据（匿名件逐字片段，PDF 注页码）与理由。"])
    ws.append(HEAD)
    for c in ws[3]:
        c.font = Font(bold=True)
    dv = DataValidation(type="list", formula1='"确认,改正"', allow_blank=True)
    dv_s = DataValidation(type="list", formula1='"stated,implied,missing,conflicting"', allow_blank=True)
    ws.add_data_validation(dv)
    ws.add_data_validation(dv_s)
    used = set()
    for name in FIELDS:
        f = rec["terms"][name]
        ws.append([name, f["status"], compact(f["value"]), spans_text(f["spans"]), field_note(rec["notes"], name),
                   "", "", "", "", ""])
        r = ws.max_row
        for col in range(6, 11):
            ws.cell(r, col).fill = YELLOW
        dv.add(ws.cell(r, 6))
        dv_s.add(ws.cell(r, 7))
        used.update(k for k in rec["notes"] if k == name or k.startswith(name + "_"))
    other = {k: v for k, v in rec["notes"].items() if k not in used and k not in ("range_max", "price_form")}
    if other:
        ws.append([])
        ws.append(["其他初标备注"])
        for k, v in other.items():
            ws.append([k, "", compact(v)])
    for i, w in enumerate(WIDTH, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    for row in ws.iter_rows(min_row=4):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


README = """# 13 个非盲标单元复核包

按《真实报价集_口径与标注规范》第八节第 1 条：这些单元先由 Claude 初标，复核人（标注人 A）对照匿名原件逐字段复核。

每个单元目录包含：
- 匿名原件（与 data/seeds/anonymous/ 中的文件相同）；
- `<seed_id>.claude_initial.json`：初标 LabelRecord；
- `<seed_id>_复核表.xlsx`：18 个条款字段各一行，黄色列由复核人填写。

复核规则：
1. 只看匿名原件、规范与 schema 字段定义；不看待评系统输出、计算器结果与其他标注人的标签。
2. 每行选“确认”或“改正”。改正时填写改正后 status、改正值（按规范第六节归一化）、原文依据（逐字片段，PDF 注页码）与理由。
3. 同一条款两处取值不同记 conflicting（规范第七节）；无法唯一确定记 missing 并写理由。
4. 复核表填完后记录用时，交回整理原件的队友导入。

初标出处已做子串核对：每个片段在去除首尾空白、连续空白合并后是匿名件正文的子串（PDF 正文取 pdftotext -layout 输出）。
"""


def main(argv) -> int:
    src, out = Path(argv[1]), Path(argv[2])
    out.mkdir(parents=True, exist_ok=True)
    (out / "README.md").write_text(README, encoding="utf-8")
    for p in sorted(src.glob("*.claude_initial.json")):
        rec = json.loads(p.read_text(encoding="utf-8"))
        sid = rec["seed_id"]
        d = out / sid
        d.mkdir(exist_ok=True)
        shutil.copy2(p, d / p.name)
        for ext in ("pdf", "txt"):
            a = ANON / f"{rec['meta']['source_id']}.{ext}"
            if a.exists():
                shutil.copy2(a, d / a.name)
        wb = Workbook()
        sheet(wb, rec)
        wb.save(d / f"{sid}_复核表.xlsx")
        print(sid)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
