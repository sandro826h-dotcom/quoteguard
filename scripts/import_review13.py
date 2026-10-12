"""回填 13 个非盲标单元的复核结果：读 output/label_review_13/ 中各单元已填的复核表（xlsx）与 Claude 初标，
生成金标签 data/seeds/labels/<seed_id>.gold.json 与复核记录 data/seeds/review13.json，并做 schema 校验与出处核对
（规范第八节第 1 条、第九节）。

用法：
  PYTHONPATH=src python3 scripts/import_review13.py --reviewer 队员乙 --date 2026-10-13 [--check-only]
      [--review-dir output/label_review_13/units] [--initial-dir output/label_review_13/labels]

复核表由 scripts/make_review_pack.py 生成，每个条款字段一行，黄色列“确认/改正”“改正后 status”“改正值”
“原文依据”“理由”由复核人填写：
- 确认：沿用初标的 status、value 与出处；原文依据如填写，替换初标出处。
- 改正：改正后 status 必填；不是 missing 时改正值与原文依据必填。改正值按 JSON 填写（数值、列表、对象），
  channel、mode、liability 直接写文字；conflicting 填两个候选组成的 JSON 列表。原文依据每行一个逐字片段，
  可写“[Q01 p1] 片段”。
单元集合为 units.json 中不在盲标名单内的单元。全部检查通过才写出文件；否则逐条列出问题，退出码 1。
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from label_common import (FIELDS, JUDGED_META, STATUSES, Docs, ImportError_, build_record, canon_value,  # noqa: E402
                          cell_text, check_anon_hashes, check_record_spans, col, dump, field_of, find_header,
                          load_records, parse_spans, parse_value, same_field, sha256_file, units_by_id)

ROOT = Path(__file__).resolve().parent.parent


def read_sheet(xlsx: Path, sid: str, problems: list) -> dict:
    ws = load_workbook(xlsx, data_only=True).worksheets[0]
    r0, heads = find_header(ws, [["字段"], ["确认/改正"]])
    if r0 is None:
        problems.append(f"{xlsx.name}: 找不到含“字段”“确认/改正”的表头行")
        return {}
    c = {k: col(heads, a) for k, a in (("name", ["字段"]), ("decision", ["确认/改正"]),
                                      ("status", ["改正后 status"]), ("value", ["改正值"]),
                                      ("evidence", ["原文依据"]), ("reason", ["理由"]))}
    lack = [k for k in ("status", "value", "evidence", "reason") if c[k] is None]
    if lack:
        problems.append(f"{xlsx.name}: 找不到列 {lack}（复核表须由 make_review_pack.py 生成，列名不要改）")
        return {}
    rows = {}
    for r in range(r0 + 1, ws.max_row + 1):
        name = cell_text(ws.cell(r, c["name"]).value)
        if name not in FIELDS:
            continue
        if name in rows:
            problems.append(f"{sid} 第 {r} 行：字段 {name} 重复")
        rows[name] = {k: (ws.cell(r, ci).value if ci else None) for k, ci in c.items()} | {"row": r}
    for name in FIELDS:
        if name not in rows:
            problems.append(f"{sid}: 复核表缺少字段 {name}")
    return rows


def run(args) -> dict:
    problems: list = []
    units = units_by_id(args.units)
    blind = set(json.loads(Path(args.blind).read_text(encoding="utf-8"))["selected_ids"])
    ids = [sid for sid in units if sid not in blind]
    initial = load_records(args.initial_dir)
    review_dir = Path(args.review_dir)
    sheets = {}
    for sid in ids:
        if sid not in initial:
            problems.append(f"{sid}: 找不到初标 JSON")
        x = review_dir / sid / f"{sid}_复核表.xlsx"
        if not x.exists():
            x = review_dir / f"{sid}_复核表.xlsx"
        if not x.exists():
            problems.append(f"{sid}: 找不到复核表 {sid}_复核表.xlsx")
            continue
        sheets[sid] = (x, read_sheet(x, sid, problems))
    if problems:
        raise ImportError_(problems)

    docs = Docs(args.anon)
    problems += check_anon_hashes(docs, units, ids)
    minutes = json.loads(Path(args.minutes_json).read_text(encoding="utf-8")) if args.minutes_json else {}
    golds, decisions, files = {}, [], {}
    for sid in ids:
        x, rows = sheets[sid]
        files[sid] = {"name": x.name, "sha256": sha256_file(x)}
        init = initial[sid]
        doc = units[sid]["source_id"]
        terms = {}
        for name in FIELDS:
            row = rows[name]
            where = f"{sid}.{name}（第 {row['row']} 行）"
            dec = "".join(cell_text(row["decision"]).split())
            base = field_of(init, name)
            ev = parse_spans(row["evidence"], doc)
            if dec == "确认":
                if cell_text(row["status"]) or cell_text(row["value"]):
                    problems.append(f"{where}: 选了“确认”却填了改正后 status 或改正值；要改值请选“改正”")
                    continue
                if base is None:
                    problems.append(f"{where}: 初标此字段为空，不能确认")
                    continue
                gold = {"status": base["status"], "value": base["value"], "spans": ev or base["spans"]}
            elif dec == "改正":
                st = cell_text(row["status"])
                val = parse_value(name, row["value"])
                if st not in STATUSES:
                    problems.append(f"{where}: 改正后 status 须为 {'/'.join(STATUSES)}，现为 {st!r}")
                    continue
                if st == "missing":
                    gold = {"status": "missing", "value": None, "spans": ev}
                elif val is None or not ev:
                    problems.append(f"{where}: 改正为 {st} 时改正值与原文依据都必填")
                    continue
                else:
                    gold = {"status": st, "value": val, "spans": ev}
                if not cell_text(row["reason"]):
                    problems.append(f"{where}: 改正时理由必填")
                    continue
            else:
                problems.append(f"{where}: 确认/改正 未填或无法识别（{cell_text(row['decision'])!r}）")
                continue
            terms[name] = gold
            decisions.append({"seed_id": sid, "field": name, "row": row["row"],
                              "decision": "confirmed" if dec == "确认" else "corrected",
                              "changed": base is None or not same_field(base, gold, name),
                              "gold_status": gold["status"], "gold_value": json.loads(canon_value(gold["value"])),
                              "evidence": ev, "reason": cell_text(row["reason"])})
        if len(terms) != len(FIELDS):
            continue
        meta = dict(units[sid])
        for k in JUDGED_META:
            if k in (init.get("meta") or {}):
                meta[k] = init["meta"][k]
        rec, bad = build_record(sid, "gold", args.date, minutes.get(sid, 0), init.get("initial_label_model"),
                                meta, terms, init.get("notes") or {})
        problems += bad
        if not bad:
            problems += check_record_spans(docs, rec)
        golds[sid] = rec
    if problems:
        raise ImportError_(problems)

    changed = [f"{x['seed_id']}.{x['field']}" for x in decisions if x["changed"]]
    review = {
        "kind": "non_blind_review",
        "spec": "docs/plan/真实报价集_口径与标注规范.md 第八节第 1 条",
        "reviewer": args.reviewer, "date": args.date,
        "inputs": {"review_sheets": files,
                   "initial": {sid: sha256_file(p) for sid in ids
                               for p in [Path(args.initial_dir) / f"{sid}.claude_initial.json"] if p.exists()},
                   "units_sha256": sha256_file(args.units), "blind_subset_sha256": sha256_file(args.blind)},
        "units": ids,
        "summary": {"fields": len(decisions),
                    "confirmed": sum(1 for x in decisions if x["decision"] == "confirmed"),
                    "corrected": sum(1 for x in decisions if x["decision"] == "corrected"),
                    "value_or_status_changed": len(changed),
                    "correction_rate": round(len(changed) / len(decisions), 4) if decisions else None,
                    "changed_fields": changed},
        "fields": decisions,
        "span_check": {"spans": sum(len(f["spans"]) for g in golds.values() for f in g["terms"].values()),
                       "failures": 0},
    }
    return {"review": review, "golds": golds}


def write(result: dict, args) -> None:
    out = Path(args.labels_dir)
    out.mkdir(parents=True, exist_ok=True)
    hashes = {}
    for sid, rec in result["golds"].items():
        p = out / f"{sid}.gold.json"
        p.write_text(dump(rec), encoding="utf-8")
        hashes[f"{sid}.gold.json"] = sha256_file(p)
    result["review"]["gold_files"] = hashes
    Path(args.record).write_text(dump(result["review"]), encoding="utf-8")


def parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--reviewer", required=True, help="复核人代号，如 队员乙")
    p.add_argument("--date", required=True, help="复核完成日期（ISO）")
    p.add_argument("--check-only", action="store_true")
    p.add_argument("--review-dir", type=Path, default=ROOT / "output/label_review_13/units")
    p.add_argument("--initial-dir", type=Path, default=ROOT / "output/label_review_13/labels")
    p.add_argument("--units", type=Path, default=ROOT / "data/seeds/units.json")
    p.add_argument("--blind", type=Path, default=ROOT / "data/seeds/blind_subset.json")
    p.add_argument("--anon", type=Path, default=ROOT / "data/seeds/anonymous")
    p.add_argument("--labels-dir", type=Path, default=ROOT / "data/seeds/labels")
    p.add_argument("--record", type=Path, default=ROOT / "data/seeds/review13.json")
    p.add_argument("--minutes-json", default=None, help="{seed_id: 复核用时分钟}，缺省记 0")
    return p


def main(argv=None) -> int:
    args = parser().parse_args(argv)
    try:
        result = run(args)
    except ImportError_ as e:
        print(f"未写出任何文件，共 {len(e.problems)} 个问题：", file=sys.stderr)
        for p in e.problems:
            print("  - " + p, file=sys.stderr)
        return 1
    s = result["review"]["summary"]
    print(f"检查通过：{len(result['golds'])} 个单元，{s['fields']} 个字段，确认 {s['confirmed']}，改正 {s['corrected']}，"
          f"取值或状态变化 {s['value_or_status_changed']}（改正率 {s['correction_rate']}），"
          f"出处 {result['review']['span_check']['spans']} 条全部核对通过")
    if args.check_only:
        return 0
    write(result, args)
    print(f"已写出 {args.record} 与 {args.labels_dir}/<seed_id>.gold.json")
    return 0


if __name__ == "__main__":
    sys.exit(main())
