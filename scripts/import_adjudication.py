"""导入盲标子集的人工裁决：读人工裁决表（xlsx）与两份初标 JSON，生成 data/seeds/adjudication.json 与
金标签 data/seeds/labels/<seed_id>.gold.json，并做 schema 校验与出处核对（规范第八节第 4 条、第九节）。

用法：
  PYTHONPATH=src python3 scripts/import_adjudication.py --xlsx 盲标子集_人工裁决表.xlsx \
      --initial-1 <初标1.json 或目录> --initial-2 <初标2.json 或目录> [--check-only]
  PYTHONPATH=src python3 scripts/import_adjudication.py --xlsx 盲标子集_人工裁决表.xlsx --inspect

裁决表各工作表按表头名称定位列（表头可在前 10 行任一行）：
- 差异裁决：两份初标不同的条款字段，每个（单元, 字段）一行。列 seed_id（或“单元”）、字段、决定、裁决 status、
  裁决 value、原文依据、理由、裁决人、日期。决定填“采用初标1”“采用初标2”（也接受 采用A/采用B、初标1/初标2）
  或“改正”。
- 一致项复核：两份初标相同的条款字段，列同上，决定填“确认”或“改正”。
- 元数据差异：page_date、price_form、date_status 等元数据，列 seed_id、字段、决定、裁决 value、理由、裁决人、日期。
- 单元备注：列 seed_id、键（备注键，如 range_max）、决定（保留、删除、改正）、裁决 value、理由。
- 说明：不读取。
取值规则：“改正”时裁决 status 必填；status 不是 missing 时原文依据必填。裁决 value 按 JSON 填写
（数值、列表、对象），channel、mode、liability 直接写文字；conflicting 填两个候选组成的 JSON 列表。
原文依据每行一个逐字片段，可写“[Q01 p1] 片段”指定文件与页码，不写时取该单元的原件。
“采用”或“确认”的行如填写了原文依据，以原文依据替换初标出处。
两份初标中 status 为空的字段视为“还没填”，不能被“采用”或“确认”。

全部检查通过才写出文件；否则逐条列出问题，退出码 1。
"""
from __future__ import annotations

import argparse
import dataclasses
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from quoteguard.schema import SeedQuote  # noqa: E402
from label_common import (FIELDS, JUDGED_META, STATUSES, Docs, ImportError_, build_record,  # noqa: E402
                          canon_value, cell_date, cell_text, check_anon_hashes, check_record_spans, col,
                          dedup_spans, dump, field_of, find_header, load_records, parse_spans, parse_value,
                          same_field, sha256_file, units_by_id)

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_MODEL = "OpenAI Codex，两个相互隔离的上下文，2026-10-11（北京时间），运行环境未提供模型 ID"

SEED = ["seed_id", "单元", "单元编号"]
NAME = ["字段", "条款字段", "元数据字段", "键", "备注键"]
DECISION = ["决定", "裁决决定", "确认/改正"]
STATUS = ["裁决 status", "裁决status", "裁决状态", "改正后 status"]
VALUE = ["裁决 value", "裁决value", "裁决值", "改正值"]
EVIDENCE = ["原文依据"]
REASON = ["理由"]
WHO = ["裁决人"]
DATE = ["日期", "裁决日期"]

PICK = {"采用初标1": 1, "采用1": 1, "采用a": 1, "初标1": 1, "a": 1, "采用初标a": 1,
        "采用初标2": 2, "采用2": 2, "采用b": 2, "初标2": 2, "b": 2, "采用初标b": 2}
CORRECT = {"改正", "改写", "修改", "自定"}
CONFIRM = {"确认", "保留"}
DROP = {"删除"}
META_LOCKED = ("seed_id", "source_id", "raw_sha256", "anon_sha256", "is_real")
META_KEYS = tuple(f.name for f in dataclasses.fields(SeedQuote))


def _dec_key(s: str) -> str:
    return "".join(s.split()).lower()


def sheet_of(wb, keyword: str):
    for ws in wb.worksheets:
        if keyword in ws.title:
            return ws
    return None


COLS = (("seed", SEED), ("name", NAME), ("decision", DECISION), ("status", STATUS), ("value", VALUE),
        ("evidence", EVIDENCE), ("reason", REASON), ("who", WHO), ("date", DATE))
# 各工作表必须找到的列（裁决人、日期在命令行给出缺省值时可以没有）
REQUIRED = {"差异裁决": ("status", "value", "evidence", "reason", "who", "date"),
            "一致项复核": ("status", "value", "evidence", "reason", "who", "date"),
            "元数据差异": ("value", "reason", "who", "date"),
            "单元备注": ("value", "reason")}
ALIASES = dict(COLS)


def read_rows(ws, sheet: str, problems: list, args=None) -> list:
    """按表头读一个工作表的非空行，返回字典列表（含行号）。必需列找不到时报错，不静默忽略。"""
    r0, heads = find_header(ws, [SEED, NAME, DECISION])
    if r0 is None:
        problems.append(f"工作表“{ws.title}”：前 10 行找不到含 seed_id/单元、字段/键、决定 的表头行")
        return []
    cols = {k: col(heads, a) for k, a in COLS}
    for k in REQUIRED.get(sheet, ()):
        if cols[k] is None and not (args is not None and k == "who" and args.adjudicator) \
                and not (args is not None and k == "date" and args.date):
            problems.append(f"工作表“{ws.title}”：找不到列 {'/'.join(ALIASES[k])}（可用 --inspect 查看表头）")
    rows = []
    for r in range(r0 + 1, ws.max_row + 1):
        get = {k: (ws.cell(r, c).value if c else None) for k, c in cols.items()}
        if not cell_text(get["seed"]) and not cell_text(get["name"]):
            continue
        get["row"] = r
        get["sheet"] = sheet
        rows.append(get)
    return rows


def inspect(xlsx: Path) -> None:
    wb = load_workbook(xlsx, data_only=True)
    for ws in wb.worksheets:
        r0, heads = find_header(ws, [SEED, NAME, DECISION])
        print(f"[{ws.title}] {ws.max_row} 行 × {ws.max_column} 列；表头行：{r0}")
        if r0:
            print("  列：" + "、".join(cell_text(c.value) for c in ws[r0] if cell_text(c.value)))
            print("  识别：" + "；".join(f"{k}→{'第 ' + str(col(heads, a)) + ' 列' if col(heads, a) else '未找到'}"
                                     for k, a in COLS))
            print(f"  数据行：{sum(1 for r in range(r0 + 1, ws.max_row + 1) if cell_text(ws.cell(r, 1).value))}")


def _who_date(row, args, problems, where):
    who = cell_text(row["who"]) or args.adjudicator
    try:
        d = cell_date(row["date"]) or args.date
    except ValueError as e:
        problems.append(f"{where}: {e}")
        d = None
    if not who:
        problems.append(f"{where}: 裁决人未填")
    if not d:
        problems.append(f"{where}: 日期未填")
    return who, d


def _hash_input(p: Path) -> str:
    p = Path(p)
    if p.is_file():
        return sha256_file(p)
    h = hashlib.sha256()
    for f in sorted(p.glob("*.json")):
        h.update(f.name.encode() + b"\0" + sha256_file(f).encode() + b"\n")
    return h.hexdigest()


def run(args) -> dict:
    problems: list = []
    units = units_by_id(args.units)
    blind = json.loads(Path(args.blind).read_text(encoding="utf-8"))
    ids = list(blind["selected_ids"])
    init = {1: load_records(args.initial_1), 2: load_records(args.initial_2)}
    for k in (1, 2):
        extra = sorted(set(init[k]) - set(ids))
        lack = [s for s in ids if s not in init[k]]
        if extra:
            problems.append(f"初标{k} 含盲标名单以外的单元：{extra}")
        if lack:
            problems.append(f"初标{k} 缺少单元：{lack}")
    for sid in ids:
        if sid not in units:
            problems.append(f"{sid}: 不在 units.json 中")
    if problems:
        raise ImportError_(problems)

    wb = load_workbook(args.xlsx, data_only=True)
    sheets = {}
    for key in ("差异裁决", "一致项复核", "元数据差异", "单元备注"):
        ws = sheet_of(wb, key)
        if ws is None:
            if key in ("差异裁决", "一致项复核"):
                problems.append(f"找不到工作表“{key}”")
            sheets[key] = []
        else:
            sheets[key] = read_rows(ws, key, problems, args)

    # ---------------- 条款字段
    term_rows: dict = {}
    for key in ("差异裁决", "一致项复核"):
        for row in sheets[key]:
            sid, name = cell_text(row["seed"]), cell_text(row["name"])
            where = f"{key} 第 {row['row']} 行（{sid}.{name}）"
            if sid not in ids:
                problems.append(f"{where}: 单元不在盲标名单中")
                continue
            if name not in FIELDS:
                problems.append(f"{where}: 未知条款字段")
                continue
            if (sid, name) in term_rows:
                problems.append(f"{where}: 与{term_rows[(sid, name)]['sheet']} 第 {term_rows[(sid, name)]['row']} 行重复")
                continue
            term_rows[(sid, name)] = row

    gold_terms: dict = {sid: {} for sid in ids}
    decisions = []
    for sid in ids:
        doc = units[sid]["source_id"]
        for name in FIELDS:
            f1, f2 = field_of(init[1][sid], name), field_of(init[2][sid], name)
            agree = f1 is not None and f2 is not None and same_field(f1, f2, name)
            row = term_rows.get((sid, name))
            if row is None:
                problems.append(f"{sid}.{name}: 裁决表中没有这一行（两份初标{'一致' if agree else '不同'}，"
                                f"应在“{'一致项复核' if agree else '差异裁决'}”）")
                continue
            where = f"{row['sheet']} 第 {row['row']} 行（{sid}.{name}）"
            if agree and row["sheet"] == "差异裁决":
                problems.append(f"{where}: 两份初标此字段相同，却列在差异裁决中")
            if not agree and row["sheet"] == "一致项复核":
                problems.append(f"{where}: 两份初标此字段不同，却列在一致项复核中")
            dec = _dec_key(cell_text(row["decision"]))
            who, d = _who_date(row, args, problems, where)
            ev = parse_spans(row["evidence"], doc)
            status_in = cell_text(row["status"])
            value_in = parse_value(name, row["value"])
            if not dec:
                problems.append(f"{where}: 决定未填")
                continue
            if dec in CORRECT:
                if status_in not in STATUSES:
                    problems.append(f"{where}: 改正时裁决 status 须为 {'/'.join(STATUSES)}，现为 {status_in!r}")
                    continue
                if status_in == "missing":
                    gold = {"status": "missing", "value": None, "spans": ev}
                else:
                    if value_in is None:
                        problems.append(f"{where}: 改正为 {status_in} 时裁决 value 必填")
                        continue
                    if not ev:
                        problems.append(f"{where}: 改正为 {status_in} 时原文依据必填")
                        continue
                    gold = {"status": status_in, "value": value_in, "spans": ev}
                kind = "corrected"
            elif dec in PICK or dec in CONFIRM:
                if dec in CONFIRM:
                    if not agree:
                        problems.append(f"{where}: 两份初标不同，决定须为采用初标1、采用初标2 或改正")
                        continue
                    base = f1
                    spans = dedup_spans(f1["spans"] + f2["spans"])
                    kind = "confirmed"
                else:
                    k = PICK[dec]
                    base = f1 if k == 1 else f2
                    if base is None:
                        problems.append(f"{where}: 初标{k} 此字段未填，不能采用")
                        continue
                    spans = list(base["spans"])
                    kind = f"adopted_{k}"
                filled = {"status": status_in or base["status"],
                          "value": base["value"] if row["value"] in (None, "") else value_in}
                if not same_field(filled, base, name):
                    problems.append(f"{where}: 决定为{cell_text(row['decision'])}，但填写的裁决 status/value 与所采用的初标不同；"
                                    "要改值请把决定改为“改正”")
                    continue
                gold = {"status": base["status"], "value": base["value"], "spans": ev or spans}
            else:
                problems.append(f"{where}: 无法识别的决定 {cell_text(row['decision'])!r}")
                continue
            gold_terms[sid][name] = gold
            decisions.append({"seed_id": sid, "field": name, "sheet": row["sheet"], "row": row["row"],
                              "decision": kind, "initials_agree": agree,
                              "gold_status": gold["status"], "gold_value": json.loads(canon_value(gold["value"])),
                              "evidence": ev, "reason": cell_text(row["reason"]), "adjudicator": who, "date": d})

    # ---------------- 元数据
    meta_rows = {}
    for row in sheets["元数据差异"]:
        sid, key = cell_text(row["seed"]), cell_text(row["name"])
        where = f"元数据差异 第 {row['row']} 行（{sid}.{key}）"
        if sid not in ids:
            problems.append(f"{where}: 单元不在盲标名单中")
        elif key not in META_KEYS:
            problems.append(f"{where}: {key!r} 不是 SeedQuote 的元数据字段")
        elif key in META_LOCKED:
            problems.append(f"{where}: {key} 以 units.json 为准，不能裁决")
        elif (sid, key) in meta_rows:
            problems.append(f"{where}: 重复")
        else:
            meta_rows[(sid, key)] = row
    gold_meta, meta_log = {}, []
    for sid in ids:
        meta = dict(units[sid])
        m1, m2 = init[1][sid].get("meta") or {}, init[2][sid].get("meta") or {}
        keys = list(JUDGED_META) + [k for (s, k) in meta_rows if s == sid and k not in JUDGED_META]
        for key in keys:
            row = meta_rows.get((sid, key))
            if row is None:
                if key in m1 or key in m2:
                    v1, v2 = m1.get(key), m2.get(key)
                    if canon_value(v1) != canon_value(v2):
                        problems.append(f"{sid}.meta.{key}: 两份初标不同（{v1!r} / {v2!r}），元数据差异表中没有这一行")
                    else:
                        meta[key] = v1
                continue
            where = f"元数据差异 第 {row['row']} 行（{sid}.{key}）"
            dec = _dec_key(cell_text(row["decision"]))
            who, d = _who_date(row, args, problems, where)
            if dec in PICK:
                v = (m1 if PICK[dec] == 1 else m2).get(key)
            elif dec in CORRECT:
                t = cell_text(row["value"])
                v = None if t.lower() in ("", "null", "none") else (cell_date(row["value"]) if key.endswith("_date")
                                                                     or key == "captured_at" else t)
            elif dec in CONFIRM:
                if canon_value(m1.get(key)) != canon_value(m2.get(key)):
                    problems.append(f"{where}: 两份初标此项不同，决定须为采用初标1、采用初标2 或改正")
                    continue
                v = m1.get(key, meta.get(key))
            else:
                problems.append(f"{where}: 无法识别的决定 {cell_text(row['decision'])!r}")
                continue
            meta[key] = v
            meta_log.append({"seed_id": sid, "key": key, "decision": cell_text(row["decision"]), "value": v,
                             "initial_1": m1.get(key), "initial_2": m2.get(key), "reason": cell_text(row["reason"]),
                             "adjudicator": who, "date": d})
        gold_meta[sid] = meta

    # ---------------- 备注
    note_rows = {}
    for row in sheets["单元备注"]:
        sid, key = cell_text(row["seed"]), cell_text(row["name"])
        where = f"单元备注 第 {row['row']} 行（{sid}.{key}）"
        if sid not in ids:
            problems.append(f"{where}: 单元不在盲标名单中")
        elif (sid, key) in note_rows:
            problems.append(f"{where}: 重复")
        else:
            note_rows[(sid, key)] = row
    gold_notes, note_log = {}, []
    for sid in ids:
        n1, n2 = init[1][sid].get("notes") or {}, init[2][sid].get("notes") or {}
        notes = {k: v for k, v in n1.items() if k in n2 and canon_value(v) == canon_value(n2[k])}
        for k in sorted((set(n1) | set(n2)) - set(notes)):
            if (sid, k) not in note_rows:
                problems.append(f"{sid}.notes.{k}: 两份初标不同或只有一份有（{n1.get(k)!r} / {n2.get(k)!r}），"
                                "单元备注表中没有这一行")
        for (s, k), row in note_rows.items():
            if s != sid:
                continue
            where = f"单元备注 第 {row['row']} 行（{sid}.{k}）"
            dec = _dec_key(cell_text(row["decision"]))
            v_in = parse_value("notes", row["value"])
            if dec in DROP:
                notes.pop(k, None)
            elif dec in CORRECT:
                if v_in is None:
                    problems.append(f"{where}: 改正时裁决 value 必填")
                    continue
                notes[k] = v_in
            elif dec in PICK:
                src = n1 if PICK[dec] == 1 else n2
                if k not in src:
                    problems.append(f"{where}: 初标{PICK[dec]} 没有这条备注")
                    continue
                notes[k] = src[k]
            elif dec in CONFIRM:
                if v_in is not None:
                    notes[k] = v_in
                elif k in n1 and k in n2 and canon_value(n1[k]) != canon_value(n2[k]):
                    problems.append(f"{where}: 两份初标此备注不同，保留时须在裁决 value 中写明内容")
                    continue
                elif k in n1 or k in n2:
                    notes[k] = n1.get(k, n2.get(k))
                else:
                    problems.append(f"{where}: 两份初标都没有这条备注，须用“改正”并填写内容")
                    continue
            else:
                problems.append(f"{where}: 无法识别的决定 {cell_text(row['decision'])!r}")
                continue
            note_log.append({"seed_id": sid, "key": k, "decision": cell_text(row["decision"]), "value": notes.get(k),
                             "reason": cell_text(row["reason"])})
        gold_notes[sid] = notes
    if problems:
        raise ImportError_(problems)

    # ---------------- 组装、schema 校验、出处核对
    docs = Docs(args.anon)
    problems += check_anon_hashes(docs, units, ids)
    minutes = json.loads(Path(args.minutes_json).read_text(encoding="utf-8")) if args.minutes_json else {}
    golds = {}
    for sid in ids:
        dates = [x["date"] for x in decisions if x["seed_id"] == sid and x["date"]]
        rec, bad = build_record(sid, "gold", max(dates) if dates else args.date, minutes.get(sid, 0),
                                args.initial_label_model, gold_meta[sid], gold_terms[sid], gold_notes[sid])
        problems += bad
        if not bad:
            problems += check_record_spans(docs, rec)
        golds[sid] = rec
    if problems:
        raise ImportError_(problems)

    # ---------------- 改正率（规范第八节；报告 5.7 节）
    rates = {}
    for k in (1, 2):
        wrong = [f"{sid}.{n}" for sid in ids for n in FIELDS
                 if field_of(init[k][sid], n) is None or not same_field(field_of(init[k][sid], n), golds[sid]["terms"][n], n)]
        total = len(ids) * len(FIELDS)
        rates[f"initial_{k}"] = {"fields": total, "corrected": len(wrong), "rate": round(len(wrong) / total, 4),
                                 "corrected_fields": wrong}
    counts: dict = {}
    for x in decisions:
        counts[x["decision"]] = counts.get(x["decision"], 0) + 1
    agree_n = sum(1 for x in decisions if x["initials_agree"])
    adjudication = {
        "kind": "blind_subset_adjudication",
        "spec": "docs/plan/真实报价集_口径与标注规范.md 第七、八节",
        "initial_label_model": args.initial_label_model,
        "inputs": {"xlsx": {"name": Path(args.xlsx).name, "sha256": sha256_file(args.xlsx)},
                   "initial_1": {"name": Path(args.initial_1).name, "sha256": _hash_input(args.initial_1)},
                   "initial_2": {"name": Path(args.initial_2).name, "sha256": _hash_input(args.initial_2)},
                   "units_sha256": sha256_file(args.units), "blind_subset_sha256": sha256_file(args.blind)},
        "units": ids,
        "adjudicators": sorted({x["adjudicator"] for x in decisions + meta_log if x.get("adjudicator")}),
        "summary": {"fields": len(decisions), "initials_agree": agree_n, "initials_differ": len(decisions) - agree_n,
                    "decisions": dict(sorted(counts.items())), "correction_rate": rates},
        "fields": decisions,
        "meta": meta_log,
        "notes": note_log,
        "span_check": {"spans": sum(len(f["spans"]) for g in golds.values() for f in g["terms"].values()),
                       "failures": 0},
    }
    return {"adjudication": adjudication, "golds": golds}


def write(result: dict, args) -> None:
    out = Path(args.labels_dir)
    out.mkdir(parents=True, exist_ok=True)
    hashes = {}
    for sid, rec in result["golds"].items():
        p = out / f"{sid}.gold.json"
        p.write_text(dump(rec), encoding="utf-8")
        hashes[sid] = sha256_file(p)
    result["adjudication"]["gold_files"] = {f"{sid}.gold.json": h for sid, h in hashes.items()}
    Path(args.adjudication).write_text(dump(result["adjudication"]), encoding="utf-8")


def parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--xlsx", required=True, type=Path)
    p.add_argument("--initial-1", type=Path)
    p.add_argument("--initial-2", type=Path)
    p.add_argument("--inspect", action="store_true", help="只列出各工作表的表头与行数")
    p.add_argument("--check-only", action="store_true", help="只检查，不写文件")
    p.add_argument("--units", type=Path, default=ROOT / "data/seeds/units.json")
    p.add_argument("--blind", type=Path, default=ROOT / "data/seeds/blind_subset.json")
    p.add_argument("--anon", type=Path, default=ROOT / "data/seeds/anonymous")
    p.add_argument("--labels-dir", type=Path, default=ROOT / "data/seeds/labels")
    p.add_argument("--adjudication", type=Path, default=ROOT / "data/seeds/adjudication.json")
    p.add_argument("--initial-label-model", default=DEFAULT_MODEL)
    p.add_argument("--adjudicator", default=None, help="行内裁决人为空时使用")
    p.add_argument("--date", default=None, help="行内日期为空时使用（ISO 日期）")
    p.add_argument("--minutes-json", default=None, help="{seed_id: 裁决用时分钟}，缺省记 0")
    return p


def main(argv=None) -> int:
    args = parser().parse_args(argv)
    if args.inspect:
        inspect(args.xlsx)
        return 0
    if not (args.initial_1 and args.initial_2):
        print("需要 --initial-1 与 --initial-2", file=sys.stderr)
        return 2
    try:
        result = run(args)
    except ImportError_ as e:
        print(f"未写出任何文件，共 {len(e.problems)} 个问题：", file=sys.stderr)
        for p in e.problems:
            print("  - " + p, file=sys.stderr)
        return 1
    s = result["adjudication"]["summary"]
    print(f"检查通过：{len(result['golds'])} 个单元，{s['fields']} 个字段（一致 {s['initials_agree']}，"
          f"不同 {s['initials_differ']}），决定 {s['decisions']}，出处 {result['adjudication']['span_check']['spans']} 条全部核对通过")
    for k, v in s["correction_rate"].items():
        print(f"  {k} 改正率 {v['corrected']}/{v['fields']} = {v['rate']}")
    if args.check_only:
        return 0
    write(result, args)
    print(f"已写出 {args.adjudication} 与 {args.labels_dir}/<seed_id>.gold.json")
    return 0


if __name__ == "__main__":
    sys.exit(main())
