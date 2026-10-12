"""盲标一致性抽检：读新队友独立填写的抽检表（盲标字段抽检表.xlsx，工作表“抽检”），与人工审定的金标签
data/seeds/labels/<seed_id>.gold.json 逐项比较，status 与 value 分别计算一致率，写入 data/seeds/spotcheck.json。

用法：
  PYTHONPATH=src python3 scripts/import_spotcheck.py --xlsx 盲标字段抽检表.xlsx --seed 20261012 \
      --checker 队员乙 --date <填写完成日期> [--check-only]

抽检表列（按表头名称定位，表头可在前 10 行任一行）：编号、单元、原件、渠道/目的区域、字段、字段含义、status、
value、币种、原文依据、备注。“字段”填条款字段名（如 min_kg_carton）；value 的写法同裁决表（JSON 或文字），
status 为 missing 时 value 留空。币种只用于价格类取值：value 中的档位、附加费或赔付条目没写 currency 时用该列补齐。
比较口径：status 完全相同；value 按 schema 解码后的规范形式相同（5 与 5.0 视为相同，省略的默认键视为已填默认值），
conflicting 比较两个候选。一致率分三项报告：status、value（全部抽检项）、value（金标签为 stated 或 implied 的项）。
原文依据与备注只记录，不参与计分。
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from label_common import (FIELDS, STATUSES, ImportError_, canon_term, cell_date, cell_text, col,  # noqa: E402
                          dump, find_header, parse_value, sha256_file)

ROOT = Path(__file__).resolve().parent.parent
COLS = (("id", ["编号"]), ("seed", ["单元", "seed_id"]), ("doc", ["原件"]), ("region", ["渠道/目的区域"]),
        ("name", ["字段"]), ("meaning", ["字段含义"]), ("status", ["status"]), ("value", ["value"]),
        ("currency", ["币种"]), ("evidence", ["原文依据"]), ("note", ["备注"]))
REQUIRED = ("id", "seed", "name", "status", "value", "currency")


def _fill_currency(v, cur: str):
    if not cur:
        return v
    if isinstance(v, list):
        return [_fill_currency(x, cur) for x in v]
    if isinstance(v, dict) and ("price" in v or "amount" in v or "basis" in v) and not v.get("currency"):
        return {**v, "currency": cur}
    return v


def read_items(xlsx: Path, problems: list) -> list:
    wb = load_workbook(xlsx, data_only=True)
    if "抽检" not in wb.sheetnames:
        problems.append(f"{xlsx.name}: 找不到工作表“抽检”（现有：{wb.sheetnames}）")
        return []
    ws = wb["抽检"]
    r0, heads = find_header(ws, [["编号"], ["字段"], ["status"]])
    if r0 is None:
        problems.append("工作表“抽检”：前 10 行找不到含 编号、字段、status 的表头行")
        return []
    c = {k: col(heads, a) for k, a in COLS}
    lack = [a[0] for k, a in COLS if k in REQUIRED and c[k] is None]
    if lack:
        problems.append(f"工作表“抽检”：找不到列 {lack}")
        return []
    items = []
    for r in range(r0 + 1, ws.max_row + 1):
        get = {k: (ws.cell(r, ci).value if ci else None) for k, ci in c.items()}
        if not cell_text(get["id"]) and not cell_text(get["seed"]):
            continue
        get["row"] = r
        items.append(get)
    return items


def run(args) -> dict:
    problems: list = []
    items = read_items(args.xlsx, problems)
    if problems:
        raise ImportError_(problems)
    blind = set(json.loads(Path(args.blind).read_text(encoding="utf-8"))["selected_ids"])
    golds, seen, results = {}, set(), []
    for it in items:
        iid, sid, name = cell_text(it["id"]), cell_text(it["seed"]), cell_text(it["name"])
        where = f"抽检 第 {it['row']} 行（{iid} {sid}.{name}）"
        if iid in seen:
            problems.append(f"{where}: 编号重复")
            continue
        seen.add(iid)
        if sid not in blind:
            problems.append(f"{where}: 单元不在盲标名单中")
            continue
        if name not in FIELDS:
            problems.append(f"{where}: 未知条款字段（应填字段名，如 min_kg_carton）")
            continue
        if sid not in golds:
            p = Path(args.labels_dir) / f"{sid}.gold.json"
            if not p.exists():
                problems.append(f"{where}: 找不到金标签 {p.name}，先运行 import_adjudication.py")
                continue
            golds[sid] = json.loads(p.read_text(encoding="utf-8"))
        st = cell_text(it["status"])
        if st not in STATUSES:
            problems.append(f"{where}: status 须为 {'/'.join(STATUSES)}，现为 {st!r}")
            continue
        val = _fill_currency(parse_value(name, it["value"]), cell_text(it["currency"]))
        if st == "missing" and val is not None:
            problems.append(f"{where}: status 为 missing 时 value 应留空")
            continue
        if st != "missing" and val is None:
            problems.append(f"{where}: status 为 {st} 时 value 必填")
            continue
        g = golds[sid]["terms"][name]
        s_ok = st == g["status"]
        v_ok = canon_term(name, val) == canon_term(name, g["value"])
        results.append({"id": iid, "seed_id": sid, "field": name, "status": st, "value": json.loads(
            canon_term(name, val)), "gold_status": g["status"], "gold_value": json.loads(canon_term(name, g["value"])),
            "status_agree": s_ok, "value_agree": v_ok, "evidence": cell_text(it["evidence"]),
            "note": cell_text(it["note"])})
    if not items:
        problems.append("工作表“抽检”没有数据行")
    if problems:
        raise ImportError_(problems)

    def rate(rows, key):
        n = len(rows)
        k = sum(1 for x in rows if x[key])
        return {"agree": k, "n": n, "rate": round(k / n, 4) if n else None}

    present = [x for x in results if x["gold_status"] in ("stated", "implied")]
    gold_files = {f"{sid}.gold.json": sha256_file(Path(args.labels_dir) / f"{sid}.gold.json") for sid in sorted(golds)}
    return {
        "kind": "blind_subset_spotcheck",
        "sampling": {"population": "盲标子集 15 个单元 × 18 个条款字段", "n": len(results), "seed": args.seed},
        "checker": args.checker, "date": args.date,
        "inputs": {"xlsx": {"name": Path(args.xlsx).name, "sha256": sha256_file(args.xlsx)},
                   "gold_files": gold_files},
        "summary": {"status": rate(results, "status_agree"), "value": rate(results, "value_agree"),
                    "value_gold_present": rate(present, "value_agree"),
                    "disagreements": [x["id"] for x in results if not (x["status_agree"] and x["value_agree"])]},
        "items": results,
    }


def parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--xlsx", required=True, type=Path)
    p.add_argument("--seed", required=True, type=int, help="抽样种子，记入结果")
    p.add_argument("--checker", required=True, help="抽检人代号")
    p.add_argument("--date", required=True, help="填写完成日期（ISO）")
    p.add_argument("--check-only", action="store_true")
    p.add_argument("--blind", type=Path, default=ROOT / "data/seeds/blind_subset.json")
    p.add_argument("--labels-dir", type=Path, default=ROOT / "data/seeds/labels")
    p.add_argument("--out", type=Path, default=ROOT / "data/seeds/spotcheck.json")
    return p


def main(argv=None) -> int:
    args = parser().parse_args(argv)
    try:
        args.date = cell_date(args.date)
        result = run(args)
    except (ImportError_, ValueError) as e:
        probs = getattr(e, "problems", [str(e)])
        print(f"未写出任何文件，共 {len(probs)} 个问题：", file=sys.stderr)
        for p in probs:
            print("  - " + p, file=sys.stderr)
        return 1
    s = result["summary"]
    print(f"抽检 {result['sampling']['n']} 项：status 一致 {s['status']['agree']}/{s['status']['n']}"
          f"（{s['status']['rate']}），value 一致 {s['value']['agree']}/{s['value']['n']}（{s['value']['rate']}），"
          f"金标签有值的项 value 一致 {s['value_gold_present']['agree']}/{s['value_gold_present']['n']}；"
          f"不一致编号 {s['disagreements']}")
    if not args.check_only:
        Path(args.out).write_text(dump(result), encoding="utf-8")
        print(f"已写出 {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
