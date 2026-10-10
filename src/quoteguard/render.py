"""Render a candidate quote as text, a spreadsheet plus markdown view, or a PNG image.

Every field that appears in the rendered quote is recorded in span_map with its location
(character range, spreadsheet cell, or pixel box). Dev and test pools use disjoint wording,
layouts and field orders. The module is deterministic: the same RenderSpec yields the same
text, the same cells and the same render_hash.

In the renderer a field with status implied is written exactly like a stated field.

Image candidates: models receive only the PNG. For image modality RenderResult.text_view is the
transcript of the drawn text and serves only as the gold transcription for scoring; it must not be
given to a model as input.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from dataclasses import dataclass, field, replace
from decimal import Decimal
from pathlib import Path
from typing import Any, Optional

from .schema import Claim, DDP, QuoteTerms, RateTier, Span, Surcharge

RENDER_VERSION = "render-2"
POOLS = ("dev", "test")
MODALITIES = ("text", "table", "image")
LOCATIONS = ("body", "footnote", "annex")

CANON = ["channel", "mode", "rate_tiers", "min_kg_shipment", "min_kg_carton", "dim_divisor",
         "rounding_kg", "surcharges", "fuel_pct", "includes", "excludes", "ddp", "claims",
         "broker_fee_cny", "isf_fee_cny", "transit_days", "validity", "liability"]
_DECIMAL_FIELDS = {"min_kg_shipment", "min_kg_carton", "dim_divisor", "rounding_kg", "fuel_pct",
                   "broker_fee_cny", "isf_fee_cny"}

# ====================================================================== lexicon
# S: chat sentence patterns. L: (table label, table value pattern).
_DEV: dict = {
    "money": {"CNY": "{a}元", "USD": "{a}美元"},
    "unit": {"per_kg": "/kg", "per_ctn": "/箱", "per_shpt": "/票"},
    "pct": "运费的{x}%",
    "mode": {"express": "国际快递", "air": "空运", "sea": "海运"},
    "join": "、",
    "rng_closed": "{lo}-{hi}kg", "rng_open": "{lo}kg起",
    "cond": {"always": "", "residential": "送住宅地址", "remote": "偏远地址",
             "ctn_longest_side_gt": "单箱最长边超{t}cm", "ctn_kg_gt": "单箱超{t}kg"},
    "cond_wrap": "（{c}）", "cont": "，按实际发生收取",
    "scope": {"duty": "仅含关税", "duty_mpf": "含关税和MPF", "all": "含全部进口环节税费"},
    "scope_wrap": "，{s}",
    "reg": "，登记号{r}", "jur": "，管辖{j}", "juris": {"US": "美国", "CN": "中国内地"},
    "ent_label": "报价主体：",
    "names": {"RES": "住宅派送费", "REM": "偏远地区附加费", "OVS": "超尺寸费", "OVW": "超重费",
              "INSP": "查验费"},
    "types": {"NVOCC": "NVOCC", "FMC_OTI": "美国 FMC 备案无船承运人",
              "MOT_NVOCC": "交通部无船承运业务备案", "AEO_CN": "海关 AEO 高级认证企业"},
    "ior_w": {"forwarder": "货代", "seller": "卖家"},
    "S": {
        "ior": "进口商（IOR）：{v}",
        "channel": "走{v}渠道", "mode": "运输方式为{v}",
        "tier": "{rng}，单价{p}",
        "min_kg_shipment": "整票{v}kg起计", "min_kg_carton": "单箱不足{v}kg按{v}kg算",
        "dim_divisor": "体积重除以{v}", "rounding_kg": "计费重按{v}kg进位",
        "surcharge": "{name}：{p}{cond}{cont}", "fuel_pct": "燃油附加费{v}%",
        "includes": "报价包含：{v}", "excludes": "不含：{v}",
        "broker_fee_cny": "清关费{p}", "isf_fee_cny": "ISF申报费{p}",
        "transit": "时效{a}-{b}天", "transit1": "时效{a}天",
        "validity": "报价有效期{a}至{b}", "liability": "赔付：{v}",
        "min_kg_shipment0": "不设最低计费重", "fuel_pct0": "不收燃油附加费",
        "isf_fee_cny0": "不收ISF申报费", "excludes0": "无其他不含费用",
        "broker_fee_cny0": "不收清关费", "min_kg_carton0": "单箱不设最低计费重", "surcharge0": "不收{name}",
    },
    "L": {
        "ior": ("进口商", "{v}"),
        "channel": ("渠道", "{v}"), "mode": ("运输方式", "{v}"),
        "min_kg_shipment": ("整票最低计费重", "{v}kg"), "min_kg_carton": ("单箱最低计费重", "{v}kg"),
        "dim_divisor": ("体积重除数", "{v}"), "rounding_kg": ("计费重进位", "{v}kg"),
        "fuel_pct": ("燃油附加费", "{v}%"), "includes": ("报价包含", "{v}"),
        "excludes": ("不含费用", "{v}"), "broker_fee_cny": ("清关费", "{p}"),
        "isf_fee_cny": ("ISF申报费", "{p}"), "transit": ("时效", "{a}-{b}天"),
        "transit1": ("时效", "{a}天"), "validity": ("有效期", "{a}至{b}"),
        "liability": ("赔付", "{v}"), "ddp": ("税务条款", "{s}"), "claims": ("资质", "{s}"),
        "surcharge": ("{name}", "{p}{cond}{cont}"), "tier": ("{rng}", "{p}"),
        "min_kg_shipment0": ("整票最低计费重", "不设"), "fuel_pct0": ("燃油附加费", "不收"),
        "isf_fee_cny0": ("ISF申报费", "不收"), "excludes0": ("不含费用", "无"),
        "broker_fee_cny0": ("清关费", "不收"), "min_kg_carton0": ("单箱最低计费重", "不设"),
        "surcharge0": ("{name}", "不收"),
    },
    "ddp_w": {
        "dev-ddp-1": {"none": "不含税，关税客户自理", "included": "双清包税",
                      "actual": "税金按实际发生报销", "capped": "包税但税金封顶{cap}"},
        "dev-ddp-2": {"none": "关税另计", "included": "含税到门",
                      "actual": "税金实报实销", "capped": "含税到门，税金超过{cap}的部分另收"},
        "dev-ddp-3": {"none": "DAP，进口环节税由买家负担", "included": "DDP 包税",
                      "actual": "DDP，税金按实际结算", "capped": "DDP，税金上限{cap}"},
        "dev-ddp-4": {"none": "到仓价，税金客户另付", "included": "税金全包，客户无需另付",
                      "actual": "进口税据实结算，多退少补", "capped": "税金包干，最高承担{cap}"},
    },
    "claim_w": {
        "dev-claim-1": "我司为{t}（{e}{rj}）",
        "dev-claim-2": "公司资质：{t}，主体{e}{rj}",
        "dev-claim-3": "持有{t}，持证公司{e}{rj}",
        "dev-claim-4": "{e}已取得{t}{rj}",
    },
    "close_mix": "FBA 头程可订, pls confirm",
}
_DEV_MIX = {"美元": " USD", "元": " CNY", "/kg": " per kg", "/箱": " per ctn", "/票": " per shpt",
            "燃油附加费": "FSC ", "偏远": "Remote area ", "住宅": "Residential ", "超尺寸": "Oversize ",
            "整票": "Whole lot ", "国际快递": "Courier", "空运": "Air", "海运": "Sea",
            "清关费": "Customs filing fee ", "时效": "Transit ", "天": " days", "您好": "Hi",
            "赔付": "Claim ", "进口商（IOR）": "Importer (IOR)", "进口商": "Importer", "货代": "Forwarder", "卖家": "Seller", "报价主体": "Quoted by"}

_TEST: dict = {
    "money": {"CNY": "RMB {a}", "USD": "US$ {a}"},
    "unit": {"per_kg": "/公斤", "per_ctn": "/件", "per_shpt": "/单"},
    "pct": "运费×{x}%",
    "mode": {"express": "快递专线", "air": "空派", "sea": "海派"},
    "join": "；",
    "rng_closed": "{lo} ~ {hi} kg", "rng_open": "{lo} kg 以上",
    "cond": {"always": "", "residential": "收货地为住宅", "remote": "目的地属偏远区域",
             "ctn_longest_side_gt": "纸箱单边长度大于{t}cm", "ctn_kg_gt": "任一箱重量大于{t}kg"},
    "cond_wrap": " 条件：{c}", "cont": " 仅在发生时收取",
    "scope": {"duty": "覆盖范围：关税", "duty_mpf": "覆盖范围：关税、MPF",
              "all": "覆盖范围：关税及全部杂税"},
    "scope_wrap": "（{s}）",
    "reg": " / 编号 {r}", "jur": " / {j}", "juris": {"US": "美方", "CN": "国内"},
    "ent_label": "出具方 ",
    "names": {"RES": "私宅投递费", "REM": "边远地区派送费", "OVS": "尺寸超限费", "OVW": "重量超限费",
              "INSP": "海关抽检费"},
    "types": {"NVOCC": "NVO 承运人资质", "FMC_OTI": "FMC OTI 登记的承运人",
              "MOT_NVOCC": "中国交通运输部备案的货运承运人", "AEO_CN": "中国海关AEO高级认证"},
    "ior_w": {"forwarder": "我司", "seller": "发货人"},
    "S": {
        "ior": "报关主体 IOR 为{v}",
        "channel": "线路：{v}", "mode": "运输类型 {v}",
        "tier": "重量 {rng} → {p}",
        "min_kg_shipment": "每票最低收费重量 {v} kg", "min_kg_carton": "每箱计费重量最少 {v} kg",
        "dim_divisor": "材积按长×宽×高÷{v}", "rounding_kg": "重量以 {v} kg 为单位向上取整",
        "surcharge": "加收{name} {p}{cond}{cont}", "fuel_pct": "燃油费率 {v} %",
        "includes": "费用已含 {v}", "excludes": "以下费用另计 {v}",
        "broker_fee_cny": "进口清关服务费 {p}", "isf_fee_cny": "ISF 录入费用 {p}",
        "transit": "预计 {a} 至 {b} 个自然日送达", "transit1": "预计 {a} 个自然日送达",
        "validity": "此价格适用 {a} 到 {b}", "liability": "货损责任：{v}",
        "min_kg_shipment0": "整票无最低收费重量要求", "fuel_pct0": "免收燃油费率",
        "isf_fee_cny0": "ISF 录入免费", "excludes0": "没有额外另计的费用",
        "broker_fee_cny0": "免收进口清关服务费", "min_kg_carton0": "每箱不限最少重量", "surcharge0": "免收{name}",
    },
    "L": {
        "ior": ("报关主体", "{v}"),
        "channel": ("线路", "{v}"), "mode": ("运输类型", "{v}"),
        "min_kg_shipment": ("最低收费重量（每票）", "{v} kg"),
        "min_kg_carton": ("最低收费重量(每箱)", "{v} kg"),
        "dim_divisor": ("材积除数", "÷{v}"), "rounding_kg": ("重量取整单位", "{v} kg"),
        "fuel_pct": ("燃油费率", "{v} %"), "includes": ("费用包含", "{v}"),
        "excludes": ("费用不含", "{v}"), "broker_fee_cny": ("进口清关费", "{p}"),
        "isf_fee_cny": ("ISF录入费", "{p}"), "transit": ("送达时间", "{a}~{b} 个自然日"),
        "transit1": ("送达时间", "{a} 个自然日"), "validity": ("价格适用期", "{a} ~ {b}"),
        "liability": ("货损责任", "{v}"), "ddp": ("关税/进口税", "{s}"),
        "claims": ("资质声明", "{s}"),
        "surcharge": ("{name}", "{p}{cond}{cont}"), "tier": ("{rng}", "{p}"),
        "min_kg_shipment0": ("最低收费重量（每票）", "无要求"), "fuel_pct0": ("燃油费率", "免收"),
        "isf_fee_cny0": ("ISF录入费", "免费"), "excludes0": ("费用不含", "没有"),
        "broker_fee_cny0": ("进口清关费", "免收"), "min_kg_carton0": ("最低收费重量（每箱）", "无要求"),
        "surcharge0": ("{name}", "免收"),
    },
    "ddp_w": {
        "test-ddp-1": {"none": "不含进口税费", "included": "到仓全包价，进口税费由我方承担",
                       "actual": "进口税费凭账单据实收取", "capped": "进口税费我方承担至{cap}，超出部分由客户补足"},
        "test-ddp-2": {"none": "不负责关税", "included": "全包含税，一口价",
                       "actual": "税费代垫，凭单结算", "capped": "代垫税费最高{cap}"},
        "test-ddp-3": {"none": "Delivered at Place, duty excluded", "included": "Delivered Duty Paid",
                       "actual": "Delivered Duty Paid, taxes at cost", "capped": "Delivered Duty Paid, tax cap {cap}"},
        "test-ddp-4": {"none": "客户自行申报并缴税", "included": "报价已含海关税款",
                       "actual": "税款以海关缴款书金额为准另收", "capped": "税款以海关缴款书为准，且不超过{cap}"},
    },
    "claim_w": {
        "test-claim-1": "资质声明：{t}，{e}{rj}",
        "test-claim-2": "{e} 具备 {t}{rj}",
        "test-claim-3": "本公司（{e}）{t}{rj}",
        "test-claim-4": "Qualification: {t} / {e}{rj}",
    },
    "close_mix": "Pls revert for booking",
}
_TEST_MIX = {"/公斤": "/KG", "/件": "/CARTON", "/单": "/BATCH", "燃油费率": "Fuel surcharge ",
             "住宅": "Home site", "私宅": "Home delivery ", "偏远区域": "Outlying zone",
             "边远地区": "Outlying zone ", "最低收费重量": "Min chargeable weight", "个自然日": "d",
             "你好": "Hi", "送达": "delivery ", "报关主体 IOR": "Declarant IOR",
             "我司": "Our company", "发货人": "Consignor", "出具方": "From"}


def _tidy(s: str) -> str:
    """Spacing for mixed Chinese and English text."""
    s = re.sub(r"([A-Za-z0-9%])([\u4e00-\u9fff])", r"\1 \2", s)
    s = re.sub(r"([\u4e00-\u9fff])([A-Za-z])", r"\1 \2", s)
    s = re.sub(r" {2,}", " ", s)
    s = re.sub(r" +([，。：；）、])", r"\1", s)
    return s.rstrip()


def _apply(obj: Any, rep: dict) -> Any:
    if isinstance(obj, str):
        new = obj
        for k in sorted(rep, key=len, reverse=True):
            new = new.replace(k, rep[k])
        return _tidy(new) if new != obj else obj
    if isinstance(obj, dict):
        return {k: _apply(v, rep) for k, v in obj.items()}
    if isinstance(obj, tuple):
        return tuple(_apply(v, rep) for v in obj)
    return obj


_RAW = {"dev": (_DEV, _DEV_MIX), "test": (_TEST, _TEST_MIX)}
LEX = {p: {False: {**raw, "_mix": False}, True: {**_apply(raw, mix), "_mix": True}} for p, (raw, mix) in _RAW.items()}
DDP_WORDING_IDS = {p: sorted(_RAW[p][0]["ddp_w"]) for p in POOLS}
CLAIM_WORDING_IDS = {p: sorted(_RAW[p][0]["claim_w"]) for p in POOLS}


# ====================================================================== templates
@dataclass(frozen=True)
class Tpl:
    id: str
    pool: str
    modality: str
    order: tuple
    # text flow
    greet: str = ""
    bullet: str = ""
    sep: str = "\n"
    end: str = ""
    fn_label: str = ""
    fn_sep: str = "；"
    annex_label: str = ""
    closing: str = ""
    # table
    roles: tuple = ("label", "value")
    title: str = ""
    sheets: tuple = ("", "")
    kv_hdr: tuple = ("", "")
    tier_hdr: tuple = ("", "")
    annex_hdr: tuple = ("", "")
    idx_hdr: str = ""
    # image
    style: str = ""
    width: int = 1100
    size: int = 24
    colors: tuple = ()
    align: str = "left"


def _ord(s: str) -> tuple:
    names = s.split()
    names += [n for n in CANON if n not in names]
    return tuple(names)


_ORD = {
    "dev-1": _ord("channel mode rate_tiers min_kg_shipment min_kg_carton dim_divisor rounding_kg fuel_pct "
                  "surcharges broker_fee_cny isf_fee_cny includes excludes ddp claims transit_days validity liability"),
    "dev-2": _ord("mode channel transit_days rate_tiers surcharges fuel_pct min_kg_shipment min_kg_carton "
                  "dim_divisor rounding_kg ddp claims includes excludes broker_fee_cny isf_fee_cny validity liability"),
    "dev-3": _ord("channel rate_tiers mode ddp surcharges fuel_pct includes excludes transit_days "
                  "min_kg_shipment min_kg_carton dim_divisor rounding_kg broker_fee_cny isf_fee_cny claims liability validity"),
    "test-1": _ord("claims ddp channel mode transit_days rate_tiers min_kg_shipment fuel_pct surcharges "
                   "dim_divisor rounding_kg min_kg_carton excludes includes broker_fee_cny isf_fee_cny liability validity"),
    "test-2": _ord("validity mode rate_tiers fuel_pct surcharges ddp broker_fee_cny isf_fee_cny channel "
                   "transit_days min_kg_carton min_kg_shipment rounding_kg dim_divisor includes excludes claims liability"),
    "test-3": _ord("rate_tiers transit_days channel mode surcharges excludes includes ddp claims "
                   "dim_divisor min_kg_shipment rounding_kg min_kg_carton fuel_pct isf_fee_cny broker_fee_cny liability validity"),
}


def _mk() -> dict:
    T: dict = {}

    def add(**kw):
        t = Tpl(**kw)
        T[t.id] = t

    # ---- dev text
    add(id="dev-text-1", pool="dev", modality="text", order=_ORD["dev-1"], greet="{alias}报价如下：",
        bullet="- ", sep="\n", fn_label="备注：", annex_label="附加费说明：", closing="有问题随时联系")
    add(id="dev-text-2", pool="dev", modality="text", order=_ORD["dev-2"],
        greet="您好，{alias}这边给您报价，请看：", bullet="", sep="，", end="。", fn_label="备注：",
        annex_label="附加费说明：", closing="以上，请查收。")
    add(id="dev-text-3", pool="dev", modality="text", order=_ORD["dev-3"], greet="{alias}：",
        bullet="{n}) ", sep="\n", fn_label="备注：", annex_label="附加费说明：", closing="")
    # ---- test text
    add(id="test-text-1", pool="test", modality="text", order=_ORD["test-1"], greet="【{alias}】运价单",
        bullet="· ", sep="\n", fn_label="注：", annex_label="附加费明细", closing="请回复确认订舱")
    add(id="test-text-2", pool="test", modality="text", order=_ORD["test-2"],
        greet="你好，下面是{alias}的报价：", bullet="", sep="；", end="。", fn_label="注：",
        annex_label="附加费明细", closing="收到请回复。")
    add(id="test-text-3", pool="test", modality="text", order=_ORD["test-3"], greet="Re: 报价 / {alias}",
        bullet="[{n}] ", sep="\n", fn_label="注：", annex_label="附加费明细", closing="")
    # ---- dev table
    for i, (roles, sheets, title) in enumerate([
            (("label", "value"), ("价格档", "附加费"), "{alias} 价格表"),
            (("idx", "label", "value"), ("报价", "附加费项目"), "{alias} 报价单"),
            (("label", "value", "note"), ("价格", "附加费表"), "{alias} 头程价格")], 1):
        add(id=f"dev-table-{i}", pool="dev", modality="table", order=_ORD[f"dev-{i}"], roles=roles,
            sheets=sheets, title=title, kv_hdr=("项目", "内容"), tier_hdr=("重量段", "单价"),
            annex_hdr=("附加费项目", "金额及条件"), idx_hdr="序号", fn_label="备注：")
    for i, (roles, sheets, title) in enumerate([
            (("label", "value"), ("运价表", "费用附录"), "运价单 · {alias}"),
            (("idx", "label", "value", "note"), ("价目", "杂费"), "{alias} / Rate Sheet"),
            (("label", "note", "value"), ("费率", "附加项"), "{alias} 费率一览")], 1):
        add(id=f"test-table-{i}", pool="test", modality="table", order=_ORD[f"test-{i}"], roles=roles,
            sheets=sheets, title=title, kv_hdr=("条款", "约定"), tier_hdr=("计费重区间", "运费"),
            annex_hdr=("费用名称", "标准"), idx_hdr="No.", fn_label="注：")
    # ---- dev image
    add(id="dev-image-1", pool="dev", modality="image", order=_ORD["dev-1"], style="bubble", width=1100,
        size=24, colors=("#EEF1F5", "#FFFFFF", "#3C4A5C", "#222222"), align="left",
        greet="{alias}", annex_label="附加费说明：", fn_label="备注：", closing="有问题随时联系",
        title="{alias} 价格表", sheets=("价格档", "附加费"), kv_hdr=("项目", "内容"),
        tier_hdr=("重量段", "单价"), annex_hdr=("附加费项目", "金额及条件"))
    add(id="dev-image-2", pool="dev", modality="image", order=_ORD["dev-2"], style="table", width=1200,
        size=24, colors=("#F7F7F5", "#DDE3EA", "#4A5560", "#1E1E1E"), align="left", roles=("label", "value"),
        greet="{alias}", annex_label="附加费说明：", fn_label="备注：", closing="有问题随时联系",
        title="{alias} 价格表", sheets=("价格档", "附加费"), kv_hdr=("项目", "内容"),
        tier_hdr=("重量段", "单价"), annex_hdr=("附加费项目", "金额及条件"))
    add(id="dev-image-3", pool="dev", modality="image", order=_ORD["dev-3"], style="card", width=1000,
        size=24, colors=("#E9ECEF", "#FFFDF5", "#5B6770", "#202020"), align="left",
        greet="{alias}：", annex_label="附加费说明：", fn_label="备注：", closing="")
    # ---- test image
    add(id="test-image-1", pool="test", modality="image", order=_ORD["test-1"], style="table", width=1300,
        size=24, colors=("#FFFFFF", "#E7E2D6", "#6B5E45", "#262626"), align="left", roles=("label", "value"),
        greet="{alias}", annex_label="附加费明细", fn_label="注：", closing="",
        title="运价单 · {alias}", sheets=("运价表", "费用附录"), kv_hdr=("条款", "约定"),
        tier_hdr=("计费重区间", "运费"), annex_hdr=("费用名称", "标准"))
    add(id="test-image-2", pool="test", modality="image", order=_ORD["test-2"], style="bubble", width=1200,
        size=24, colors=("#F5F1EA", "#DCE6F2", "#52606D", "#1F2933"), align="right",
        greet="{alias}", annex_label="附加费明细", fn_label="注：", closing="收到请回复。")
    add(id="test-image-3", pool="test", modality="image", order=_ORD["test-3"], style="card", width=1100,
        size=24, colors=("#FAFAFA", "#EEF3EE", "#6C757D", "#212121"), align="left",
        greet="Re: 报价 / {alias}", annex_label="附加费明细", fn_label="注：", closing="")
    return T


TEMPLATES: dict = _mk()
TEMPLATE_IDS = {p: {m: sorted(t for t, d in TEMPLATES.items() if d.pool == p and d.modality == m)
                    for m in MODALITIES} for p in POOLS}

_TPL_STR_FIELDS = ("greet", "bullet", "sep", "end", "fn_label", "fn_sep", "annex_label", "closing",
                   "title", "idx_hdr")


def sentence_patterns(pool: str) -> set:
    """All wording strings (sentence patterns, labels, headers, template phrases) used by a pool.

    Pure placeholders and pure punctuation are left out. Dev and test sets are disjoint.
    """
    out: set = set()

    def walk(o):
        if isinstance(o, str):
            if any(c.isalnum() for c in re.sub(r"\{[^}]*\}", "", o)):
                out.add(o)
        elif isinstance(o, dict):
            for v in o.values():
                walk(v)
        elif isinstance(o, (tuple, list)):
            for v in o:
                walk(v)

    for mix in (False, True):
        walk(LEX[pool][mix])
    for t in TEMPLATES.values():
        if t.pool != pool:
            continue
        for f in _TPL_STR_FIELDS:
            walk(getattr(t, f))
        for f in ("sheets", "kv_hdr", "tier_hdr", "annex_hdr"):
            walk(getattr(t, f))
            walk(_apply(getattr(t, f), _RAW[pool][1]))
        for f in _TPL_STR_FIELDS:
            walk(_apply(getattr(t, f), _RAW[pool][1]))
    return out


# ====================================================================== spec / result
def _plain(o: Any) -> Any:
    if isinstance(o, Decimal):
        return str(o)
    if hasattr(o, "to_dict"):
        return _plain(o.to_dict())
    if isinstance(o, dict):
        return {str(k): _plain(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [_plain(v) for v in o]
    return o


TAXES = ("关税", "MPF", "HMF")
_SCOPE_TAXES = {"duty": {"关税"}, "duty_mpf": {"关税", "MPF"}, "all": {"关税", "MPF", "HMF"}}


def _tax_set(lst: Any) -> set:
    return {t for t in TAXES if any(t in str(x) for x in lst)}


@dataclass
class RenderSpec:
    cand_id: str
    terms: Any
    modality: str
    template_id: str
    pool: str
    conflicts: dict = field(default_factory=dict)
    fee_location: dict = field(default_factory=dict)
    lang_mix: bool = False
    ddp_wording_id: str = ""
    claim_wording_id: str = ""
    forwarder_alias: str = "货代A"
    quote_entity: str = ""
    _conf_raw: dict = field(default_factory=dict, repr=False, compare=False)

    def __post_init__(self):
        if isinstance(self.terms, dict):
            d = {}
            for n, f in self.terms.items():
                if isinstance(f, dict) and f.get("status") == "conflicting":
                    self._conf_raw[n] = f.get("value")
                    f = {**f, "value": None}
                d[n] = f
            self.terms = QuoteTerms.from_dict(d)
        elif isinstance(self.terms, QuoteTerms):
            for n in CANON:
                f = getattr(self.terms, n)
                if f.status == "conflicting" and f.value is not None:
                    self._conf_raw[n] = f.value
        if not self.ddp_wording_id:
            self.ddp_wording_id = f"{self.pool}-ddp-1"
        if not self.claim_wording_id:
            self.claim_wording_id = f"{self.pool}-claim-1"

    def validate(self) -> None:
        if self.pool not in POOLS:
            raise ValueError(f"pool {self.pool!r}")
        if self.modality not in MODALITIES:
            raise ValueError(f"modality {self.modality!r}")
        t = TEMPLATES.get(self.template_id)
        if t is None or t.pool != self.pool or t.modality != self.modality:
            raise ValueError(f"template {self.template_id!r} does not match pool/modality")
        if self.ddp_wording_id not in LEX[self.pool][False]["ddp_w"]:
            raise ValueError(f"ddp_wording_id {self.ddp_wording_id!r} not in pool {self.pool}")
        if self.claim_wording_id not in LEX[self.pool][False]["claim_w"]:
            raise ValueError(f"claim_wording_id {self.claim_wording_id!r} not in pool {self.pool}")
        if not isinstance(self.terms, QuoteTerms):
            raise ValueError("terms must be QuoteTerms or dict")
        if not self.quote_entity or not isinstance(self.quote_entity, str):
            raise ValueError("quote_entity is required")
        for k, v in self.fee_location.items():
            if k not in CANON:
                raise ValueError(f"fee_location key {k!r} is not a quote field")
            if v not in LOCATIONS:
                raise ValueError(f"fee_location[{k!r}]={v!r}")
        for k, v in self.conflicts.items():
            if k not in CANON:
                raise ValueError(f"conflicts key {k!r} is not a quote field")
            if getattr(self.terms, k).status != "conflicting":
                raise ValueError(f"conflicts[{k!r}] given but terms.{k} is not conflicting")
            if not isinstance(v, (list, tuple)) or len(v) != 2:
                raise ValueError(f"conflicts[{k!r}] needs two values")
            if _plain(_coerce(k, v[0])) == _plain(_coerce(k, v[1])):
                raise ValueError(f"conflicts[{k!r}] holds two equal values")
            raw = self._conf_raw.get(k)
            if raw is not None:
                if (not isinstance(raw, (list, tuple)) or len(raw) != 2 or
                        [_plain(_coerce(k, x)) for x in raw] != [_plain(_coerce(k, x)) for x in v]):
                    raise ValueError(f"conflicts[{k!r}] differs from the two candidates in terms")
        for n in CANON:
            if getattr(self.terms, n).status == "conflicting" and n not in self.conflicts:
                raise ValueError(f"terms.{n} is conflicting but conflicts has no entry")
        self._check_taxes()

    def _check_taxes(self) -> None:
        t = self.terms
        inc, exc, ddp = t.includes, t.excludes, t.ddp
        known = ddp.status in ("stated", "implied") and isinstance(ddp.value, DDP)
        for nm, f in (("includes", inc), ("excludes", exc)):
            if f.status not in ("stated", "implied") or f.value is None:
                continue
            tx = _tax_set(f.value)
            if not known:
                if tx:
                    raise ValueError(f"{nm} lists taxes {sorted(tx)} but ddp is missing or conflicting")
                continue
            d = ddp.value
            want = _SCOPE_TAXES[d.scope] if d.is_ddp else set()
            if nm == "includes" and tx != want:
                raise ValueError(f"includes taxes {sorted(tx)} do not match ddp.scope (expected {sorted(want)})")
            if nm == "excludes":
                if not d.is_ddp and tx != set(TAXES):
                    raise ValueError("excludes must list 关税, MPF and HMF when the quote is not DDP")
                if d.is_ddp and tx & want:
                    raise ValueError("excludes lists a tax that ddp.scope includes")

    def to_dict(self) -> dict:
        return {"cand_id": self.cand_id, "terms": _plain(self.terms), "modality": self.modality,
                "template_id": self.template_id, "pool": self.pool, "conflicts": _plain(self.conflicts),
                "fee_location": dict(self.fee_location), "lang_mix": self.lang_mix,
                "ddp_wording_id": self.ddp_wording_id, "claim_wording_id": self.claim_wording_id,
                "forwarder_alias": self.forwarder_alias, "quote_entity": self.quote_entity}

    @classmethod
    def from_dict(cls, d: dict, **overrides: Any) -> "RenderSpec":
        """Read a Candidate.presentation dict. cand_id and terms may come from overrides."""
        m = {**d, **overrides}
        for k in ("cand_id", "terms", "modality", "template_id", "pool"):
            if k not in m:
                raise ValueError(f"RenderSpec.from_dict: missing {k}")
        pool = m["pool"]
        return cls(m["cand_id"], m["terms"], m["modality"], m["template_id"], pool,
                   dict(m.get("conflicts") or {}), dict(m.get("fee_location") or {}),
                   bool(m.get("lang_mix", False)), m.get("ddp_wording_id") or "",
                   m.get("claim_wording_id") or "", m.get("forwarder_alias") or "货代A",
                   m.get("quote_entity") or "")


@dataclass
class RenderResult:
    files: list
    text_view: str
    span_map: dict
    render_hash: str
    meta: dict = field(default_factory=dict)


# ====================================================================== items
def num(x: Any) -> str:
    d = x if isinstance(x, Decimal) else Decimal(str(x))
    return format(d.normalize(), "f")


def _pct(x: Any) -> str:
    return num(Decimal(str(x)) * 100)


@dataclass
class _Item:
    name: str
    keys: tuple
    label: str
    value: str
    sentence: str
    loc: str = "body"
    seq: int = 0


def _coerce(name: str, v: Any) -> Any:
    if name == "rate_tiers":
        return [x if isinstance(x, RateTier) else RateTier.from_dict(x) for x in v]
    if name == "surcharges":
        return [x if isinstance(x, Surcharge) else Surcharge.from_dict(x) for x in v]
    if name == "claims":
        return [x if isinstance(x, Claim) else Claim.from_dict(x) for x in v]
    if name == "ddp":
        if isinstance(v, DDP):
            return v
        if isinstance(v, bool):
            return DDP(v, "all", "included" if v else None)
        return DDP.from_dict(v)
    if name in _DECIMAL_FIELDS:
        return Decimal(str(v))
    return v


def _money(lx: dict, amount: Any, cur: str) -> str:
    return lx["money"][cur].format(a=num(amount))


def _price(lx: dict, amount: Any, cur: str, basis: str) -> str:
    return _money(lx, amount, cur) + lx["unit"][basis]


def _fin(lx: dict, s: str) -> str:
    s = s.strip()
    return _tidy(s) if lx["_mix"] else s


def _lookup(lx: dict, table: str, key: str) -> str:
    try:
        return lx[table][key]
    except KeyError:
        raise ValueError(f"no wording for {table} {key!r}") from None


def _parts(name: str, v: Any, lx: dict, spec: RenderSpec, force0: bool = False) -> list:
    """Return dicts (keys, label, value, sentence) for one field value; zero optional values give nothing."""
    S, L = lx["S"], lx["L"]
    out = []
    mode = spec.terms.mode.value if spec.terms.mode.status != "missing" else None

    def one(key, kw, keys=None):
        lab, vp = L[key]
        out.append({"keys": keys or (name,), "label": _fin(lx, lab.format(**kw)),
                    "value": _fin(lx, vp.format(**kw)), "sentence": _fin(lx, S[key].format(**kw))})

    if name == "rate_tiers":
        for t in v:
            rng = (lx["rng_closed"].format(lo=num(t.min_kg), hi=num(t.max_kg)) if t.max_kg is not None
                   else lx["rng_open"].format(lo=num(t.min_kg)))
            one("tier", {"rng": rng, "p": _price(lx, t.price, t.currency, "per_kg")})
    elif name == "surcharges":
        for s in v:
            if s.amount != 0:
                out.append(_surcharge_part(s, lx))
    elif name == "claims":
        cw = lx["claim_w"][spec.claim_wording_id]
        for c in v:
            rj = (lx["reg"].format(r=c.reg_no) if c.reg_no else "") + \
                 (lx["jur"].format(j=_lookup(lx, "juris", c.jurisdiction)) if c.jurisdiction else "")
            s = _fin(lx, cw.format(t=_lookup(lx, "types", c.type), e=c.entity, rj=rj))
            out.append({"keys": ("claims",), "label": L["claims"][0], "value": s, "sentence": s})
    elif name == "ddp":
        w = lx["ddp_w"][spec.ddp_wording_id]
        cls = "none" if not v.is_ddp else (v.duty_terms if v.duty_terms in ("included", "actual", "capped")
                                           else "included")
        cap = _price(lx, v.cap_cny_per_kg, "CNY", "per_kg") if v.cap_cny_per_kg is not None else ""
        s = w[cls].format(cap=cap)
        if v.is_ddp:
            s += lx["scope_wrap"].format(s=lx["scope"][v.scope])
        s = _fin(lx, s)
        out.append({"keys": ("ddp",), "label": L["ddp"][0], "value": s, "sentence": s})
        if v.ior in lx["ior_w"]:
            one("ior", {"v": lx["ior_w"][v.ior]}, keys=("ddp.ior",))
            out[-1]["name"] = "ddp"
            out[-1]["loc_keys"] = ("ddp",)
    elif name == "channel":
        one(name, {"v": str(v)})
    elif name == "mode":
        one(name, {"v": lx["mode"][v]})
    elif name == "min_kg_shipment":
        if v == 0:
            one("min_kg_shipment0", {})   # required term: written in words
        else:
            one(name, {"v": num(v)})
    elif name in ("min_kg_carton", "dim_divisor", "rounding_kg"):
        if name == "min_kg_carton" and v == 0:
            if force0:
                one("min_kg_carton0", {})
        else:
            one(name, {"v": num(v)})
    elif name == "fuel_pct":
        if v != 0:
            one(name, {"v": _pct(v)})
        elif mode == "express" or force0:
            one("fuel_pct0", {})
    elif name in ("includes", "excludes"):
        if v:
            one(name, {"v": lx["join"].join(v)})
        elif name == "excludes":
            one("excludes0", {})
    elif name == "broker_fee_cny":
        if v != 0:
            one(name, {"p": _price(lx, v, "CNY", "per_shpt")})
        elif force0:
            one("broker_fee_cny0", {})
    elif name == "isf_fee_cny":
        if v != 0:
            one(name, {"p": _price(lx, v, "CNY", "per_shpt")})
        elif mode == "sea" or force0:
            one("isf_fee_cny0", {})
    elif name == "transit_days":
        a, b = v
        if a == b:
            one("transit1", {"a": a}, keys=("transit_days",))
        else:
            one("transit", {"a": a, "b": b}, keys=("transit_days",))
    elif name == "validity":
        one(name, {"a": v[0], "b": v[1]})
    elif name == "liability":
        one(name, {"v": str(v)})
    else:
        raise ValueError(f"cannot render field {name}")
    return out


def _surcharge_part(s: Surcharge, lx: dict) -> dict:
    name = _lookup(lx, "names", s.code)
    if s.amount == 0:
        lab, vp = lx["L"]["surcharge0"]
        return {"keys": ("surcharges", f"surcharges.{s.code}"), "label": _fin(lx, lab.format(name=name)),
                "value": _fin(lx, vp), "sentence": _fin(lx, lx["S"]["surcharge0"].format(name=name))}
    if s.basis == "pct_freight":
        p = lx["pct"].format(x=_pct(s.amount))
    else:
        p = _price(lx, s.amount, s.currency or "CNY", s.basis)
    c = lx["cond"][s.condition].format(t=num(s.threshold) if s.threshold is not None else "")
    cond = lx["cond_wrap"].format(c=c) if c else ""
    cont = lx["cont"] if s.contingent else ""
    kw = {"name": name, "p": p, "cond": cond, "cont": cont}
    lab, vp = lx["L"]["surcharge"]
    return {"keys": ("surcharges", f"surcharges.{s.code}"), "label": _fin(lx, lab.format(**kw)),
            "value": _fin(lx, vp.format(**kw)), "sentence": _fin(lx, lx["S"]["surcharge"].format(**kw))}


def _loc(spec: RenderSpec, keys: tuple, flip: bool) -> str:
    base = "body"
    for k in reversed(keys):
        if k in spec.fee_location:
            base = spec.fee_location[k]
            break
    if not flip:
        return base
    return "footnote" if base == "body" else "body"


def _build_items(spec: RenderSpec, tpl: Tpl) -> list:
    lx = LEX[spec.pool][spec.lang_mix]
    terms: QuoteTerms = spec.terms
    conf = spec.conflicts
    items: list = []

    def push(parts, flip):
        for p in parts:
            keys = tuple(p["keys"])
            items.append(_Item(p.get("name", keys[0]), keys, p["label"], p["value"], p["sentence"],
                               _loc(spec, tuple(p.get("loc_keys", keys)), flip), len(items)))

    for name in CANON:
        f = getattr(terms, name)
        if name in conf:
            vals = [_coerce(name, v) for v in conf[name]]
            if name == "surcharges":
                a, b = vals
                by_a = {s.code: s for s in a}
                by_b = {s.code: s for s in b}
                # a zero amount is written in words when the other table gives a different value
                push([_surcharge_part(s, lx) for s in a
                      if s.amount != 0 or (s.code in by_b and by_b[s.code] != s)], False)
                # the second table repeats only the lines whose values differ
                push([_surcharge_part(s, lx) for s in b if by_a.get(s.code) != s
                      and (s.amount != 0 or s.code in by_a)], True)
            else:
                for idx, v in enumerate(vals):
                    push(_parts(name, v, lx, spec, True), idx == 1)
            continue
        if f.status == "missing" or f.value is None:
            continue
        push(_parts(name, _coerce(name, f.value), lx, spec), False)
    pos = {n: i for i, n in enumerate(tpl.order)}
    items.sort(key=lambda it: (pos.get(it.name, 99), it.seq))
    return items


def _ent_item(spec: RenderSpec) -> "_Item":
    return _Item("quote_entity", ("quote_entity",), "", "", spec.quote_entity, "body", -1)


def _m(spec: RenderSpec, s: str) -> str:
    return _apply(s, _RAW[spec.pool][1]) if spec.lang_mix else s


def _by_loc(items: list) -> dict:
    d = {k: [] for k in LOCATIONS}
    for it in items:
        d[it.loc].append(it)
    return d


def _add_span(span_map: dict, it: _Item, span: Span) -> None:
    for k in it.keys:
        span_map.setdefault(k, []).append(span)


def _safe(cand_id: str) -> str:
    return re.sub(r"[^0-9A-Za-z_.\-]", "_", cand_id)


# ====================================================================== text
def _render_text(spec: RenderSpec, tpl: Tpl, items: list, outdir: Path) -> tuple:
    name = _safe(spec.cand_id) + ".txt"
    lx = LEX[spec.pool][spec.lang_mix]
    parts: list = []
    pos = 0
    span_map: dict = {}

    def put(s: str) -> int:
        nonlocal pos
        start = pos
        parts.append(s)
        pos += len(s)
        return start

    def group(its: list, sep: str, end: str, bullet: str) -> None:
        for i, it in enumerate(its):
            if i:
                put(sep)
            put(bullet.format(n=i + 1))
            a = put(it.sentence)
            _add_span(span_map, it, Span(name, "text", it.sentence, start=a, end=a + len(it.sentence)))
        put(end)

    g = _by_loc(items)
    lines_mode = tpl.sep.startswith("\n")
    put(_m(spec, tpl.greet).format(alias=spec.forwarder_alias) + "\n")
    put(lx["ent_label"])
    a = put(spec.quote_entity)
    _add_span(span_map, _ent_item(spec), Span(name, "text", spec.quote_entity, start=a, end=a + len(spec.quote_entity)))
    put("\n")
    group(g["body"], tpl.sep, tpl.end + "\n", _m(spec, tpl.bullet))
    if g["footnote"]:
        put(_m(spec, tpl.fn_label))
        group(g["footnote"], tpl.fn_sep, "。\n", "")
    if g["annex"]:
        put(_m(spec, tpl.annex_label) + ("\n" if lines_mode else ""))
        group(g["annex"], tpl.sep, tpl.end + "\n", _m(spec, tpl.bullet))
    close = _closing(spec, tpl)
    if close:
        put(close + "\n")
    text = "".join(parts)
    (outdir / name).write_text(text, encoding="utf-8")
    return [str(outdir / name)], text, span_map, {"text": text}, {}


def _closing(spec: RenderSpec, tpl: Tpl) -> str:
    close = _m(spec, tpl.closing)
    if spec.lang_mix:
        extra = LEX[spec.pool][True]["close_mix"]
        close = (close + " " + extra) if close else extra
    return close


# ====================================================================== table model
@dataclass
class _Row:
    kind: str                 # title | meta | hdr | data | note
    cells: list
    link: Optional[int] = None  # column carrying the item span
    item: Optional[_Item] = None


@dataclass
class _Sheet:
    name: str
    rows: list


def _role_cells(roles: tuple, label: str, value: str, n: int) -> list:
    d = {"idx": str(n), "label": label, "value": value, "note": ""}
    return [d[r] for r in roles]


def _hdr_cells(tpl: Tpl, roles: tuple, pair: tuple, spec: RenderSpec) -> list:
    d = {"idx": _m(spec, tpl.idx_hdr), "label": _m(spec, pair[0]), "value": _m(spec, pair[1]), "note": ""}
    return [d[r] for r in roles]


def _table_model(spec: RenderSpec, tpl: Tpl, items: list, roles: tuple) -> list:
    g = _by_loc(items)
    lx = LEX[spec.pool][spec.lang_mix]
    vcol = roles.index("value")
    main = _Sheet(_m(spec, tpl.sheets[0]), [
        _Row("title", [_m(spec, tpl.title).format(alias=spec.forwarder_alias)]),
        _Row("meta", [lx["ent_label"] + spec.quote_entity], 0, _ent_item(spec))])
    seen: set = set()
    cur = None
    n = 0
    for it in g["body"]:
        want = "tier" if it.keys == ("rate_tiers",) else "kv"
        if want != cur:
            if want not in seen:
                main.rows.append(_Row("hdr", _hdr_cells(tpl, roles, tpl.tier_hdr if want == "tier" else tpl.kv_hdr, spec)))
                seen.add(want)
            cur = want
        n += 1
        main.rows.append(_Row("data", _role_cells(roles, it.label, it.value, n), vcol, it))
    for it in g["footnote"]:
        main.rows.append(_Row("note", [_m(spec, tpl.fn_label) + it.sentence], 0, it))
    sheets = [main]
    if g["annex"]:
        ann = _Sheet(_m(spec, tpl.sheets[1]), [_Row("hdr", _hdr_cells(tpl, roles, tpl.annex_hdr, spec))])
        for i, it in enumerate(g["annex"], 1):
            ann.rows.append(_Row("data", _role_cells(roles, it.label, it.value, i), vcol, it))
        sheets.append(ann)
    return sheets


def _md_view(sheets: list) -> tuple:
    """Markdown text plus spans (list of (item, start, end))."""
    parts: list = []
    pos = 0
    found: list = []

    def put(s):
        nonlocal pos
        a = pos
        parts.append(s)
        pos += len(s)
        return a

    for si, sh in enumerate(sheets):
        put(("\n" if si else "") + f"## {sh.name}\n\n")
        prev = None
        for r in sh.rows:
            if r.kind == "title":
                put(f"**{r.cells[0]}**\n\n")
            elif r.kind == "meta":
                s0 = put(r.cells[0])
                found.append((r.item, s0, s0 + len(r.cells[0])))
                put("\n\n")
            elif r.kind == "hdr":
                if prev in ("data", "note"):
                    put("\n")
                put("| " + " | ".join(r.cells) + " |\n")
                put("|" + "---|" * len(r.cells) + "\n")
            elif r.kind == "data":
                put("| ")
                for j, c in enumerate(r.cells):
                    s0 = put(c)
                    if j == r.link:
                        found.append((r.item, s0, s0 + len(c)))
                    put(" | " if j < len(r.cells) - 1 else " |\n")
            else:
                if prev != "note":
                    put("\n")
                put("> ")
                s0 = put(r.cells[0])
                found.append((r.item, s0, s0 + len(r.cells[0])))
                put("\n")
            prev = r.kind
    return "".join(parts), found


_ZIP_TIME = (2026, 10, 1, 0, 0, 0)


def _write_xlsx(path: Path, sheets: list) -> dict:
    from datetime import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "synthetic"
    wb.properties.lastModifiedBy = "synthetic"
    wb.properties.created = datetime(2026, 10, 1)
    wb.properties.modified = datetime(2026, 10, 1)
    dump: dict = {}
    cells: list = []
    for sh in sheets:
        ws = wb.create_sheet(sh.name)
        dump[sh.name] = [list(r.cells) for r in sh.rows]
        for ri, r in enumerate(sh.rows, 1):
            for ci, c in enumerate(r.cells, 1):
                ws.cell(row=ri, column=ci, value=c)
                if r.kind in ("title", "hdr"):
                    ws.cell(row=ri, column=ci).font = Font(bold=True)
            if r.item is not None:
                col = r.link or 0
                cells.append((sh.name, f"{get_column_letter(col + 1)}{ri}", r.cells[col], r.item))
        ncol = max(len(r.cells) for r in sh.rows)
        for ci in range(1, ncol + 1):
            ws.column_dimensions[get_column_letter(ci)].width = 34 if ci > 1 else 26
    buf = io.BytesIO()
    wb.save(buf)
    # fixed zip entry timestamps make the bytes reproducible
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for zi in src.infolist():
            ni = zipfile.ZipInfo(zi.filename, _ZIP_TIME)
            ni.compress_type = zipfile.ZIP_DEFLATED
            ni.external_attr = zi.external_attr
            data = src.read(zi.filename)
            if zi.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-10-01T00:00:00Z", data)
                data = re.sub(rb"(<dcterms:created[^>]*>)[^<]*", rb"\g<1>2026-10-01T00:00:00Z", data)
            dst.writestr(ni, data)
    return {"dump": dump, "cells": cells}


def _render_table(spec: RenderSpec, tpl: Tpl, items: list, outdir: Path) -> tuple:
    base = _safe(spec.cand_id)
    xname, mname = base + ".xlsx", base + ".md"
    sheets = _table_model(spec, tpl, items, tpl.roles)
    md, found = _md_view(sheets)
    span_map: dict = {}
    for it, a, b in found:
        _add_span(span_map, it, Span(mname, "text", md[a:b], start=a, end=b))
    info = _write_xlsx(outdir / xname, sheets)
    for sheet, cell, quote, it in info["cells"]:
        _add_span(span_map, it, Span(xname, "cell", quote, sheet=sheet, cell=cell))
    (outdir / mname).write_text(md, encoding="utf-8")
    return [str(outdir / xname), str(outdir / mname)], md, span_map, {"sheets": info["dump"], "md": md}, {}


# ====================================================================== image
_FONT_FAMILY = "Noto Sans CJK SC"
_WM = "合成示例"
MIN_FONT_PX = 18
MIN_NOTE_PX = 16
MAX_EDGE = 1600
_M: Optional["_Measure"] = None


def _font_info() -> dict:
    from matplotlib import font_manager as fm
    import matplotlib
    try:
        path = fm.findfont(fm.FontProperties(family=_FONT_FAMILY), fallback_to_default=False)
    except ValueError:
        raise RuntimeError(f"font {_FONT_FAMILY!r} not found; install Noto Sans CJK") from None
    return {"family": _FONT_FAMILY, "path": str(path), "matplotlib": matplotlib.__version__}


class _Measure:
    def __init__(self):
        from matplotlib.backends.backend_agg import RendererAgg
        self.info = _font_info()
        self.r = RendererAgg(100, 100, 100)
        self.cache: dict = {}

    def prop(self, size_px: int, bold: bool = False):
        from matplotlib.font_manager import FontProperties
        return FontProperties(family=self.info["family"], size=size_px * 72 / 100,
                              weight="bold" if bold else "normal")

    def width(self, s: str, size: int, bold: bool = False) -> float:
        k = (s, size, bold)
        if k not in self.cache:
            w, _, _ = self.r.get_text_width_height_descent(s, self.prop(size, bold), False)
            self.cache[k] = float(w)
        return self.cache[k]

    def wrap(self, s: str, size: int, maxw: float, bold: bool = False) -> list:
        # an ASCII run (MPF, ONT8, per kg, 2026-10-12) is one unit and is not split unless wider than a line
        toks = re.findall(r"[A-Za-z0-9%./\-]+(?: [A-Za-z0-9%./\-]+)*|.", s, flags=re.S)
        lines, cur, w = [], "", 0.0
        for tok in toks:
            tw = self.width(tok, size, bold)
            units = list(tok) if tw > maxw else [tok]
            for u in units:
                uw = tw if len(units) == 1 else self.width(u, size, bold)
                if cur and w + uw > maxw:
                    lines.append(cur)
                    cur, w = "", 0.0
                cur += u
                w += uw
        if cur:
            lines.append(cur)
        return lines or [""]


def _measure() -> _Measure:
    global _M
    if _M is None:
        _M = _Measure()
    return _M


class _Canvas:
    def __init__(self, W: int):
        self.W = W
        self.H = 0
        self.ops: list = []
        self.boxes: list = []   # (item, bbox, quote)
        self.texts: list = []

    def rect(self, x, y, w, h, fc, ec=None, r=0):
        self.ops.append(("rect", x, y, w, h, fc, ec, r))

    def line(self, x0, y0, x1, y1, color):
        self.ops.append(("line", x0, y0, x1, y1, color))

    def text(self, x, y, s, size, color, bold=False):
        self.ops.append(("text", x, y, s, size, color, bold))
        self.texts.append(s)


def _layout(spec: RenderSpec, tpl: Tpl, items: list, size: int, M: _Measure, limit: Optional[int]) -> tuple:
    """Lay out the quote; pages break when limit (pixels) would be exceeded. -> (canvases, transcript)."""
    W = tpl.width
    bg, fill, accent, ink = tpl.colors
    lx = LEX[spec.pool][spec.lang_mix]
    lh = int(size * 1.45)
    mar, pad, bar = 36, 12, 64
    small = max(size - 4, MIN_NOTE_PX)
    sl = int(small * 1.45)
    g = _by_loc(items)
    alias = _m(spec, tpl.greet).format(alias=spec.forwarder_alias)
    ent_item = _ent_item(spec)
    ent_text = lx["ent_label"] + spec.quote_entity
    transcript: list = [alias]
    pages: list = []
    st: dict = {}

    def new_page():
        cv = _Canvas(W)
        cv.rect(0, 0, W, bar, accent)
        cv.text(mar, (bar - size * 1.3) / 2, alias, size, "#FFFFFF", True)
        st["cv"] = cv
        if tpl.style == "card":
            st["card_top"] = bar + 14
            st["y"] = bar + 14 + 22
        elif tpl.style == "table":
            st["y"] = bar + 18
        else:
            st["y"] = bar + 20
        st["start"] = st["y"]

    def end_page():
        cv, y = st["cv"], st["y"]
        if tpl.style == "card":
            y += 12
            cv.ops.insert(1, ("rect", mar - 6, st["card_top"], W - 2 * mar + 12, y - st["card_top"], fill, "#C4CCD6", 16))
        cv.H = int(y + 56)
        cv.text(W - 24 - M.width(_WM, 18), cv.H - 38, _WM, 18, "#888888")
        pages.append(cv)

    def ensure(h):
        extra = 12 if tpl.style == "card" else 0      # card bottom padding added by end_page
        if limit is not None and st["y"] + h + 56 + extra > limit and st["y"] > st["start"]:
            end_page()
            new_page()

    def text_lines(x, lns, sz, lhh, color, bold=False):
        cv, y = st["cv"], st["y"]
        for k, l in enumerate(lns):
            cv.text(x, y + k * lhh, l, sz, color, bold)

    def note(s, it, sz, lhh):
        lns = M.wrap(s, sz, W - 2 * mar)
        h = len(lns) * lhh
        ensure(h + 4)
        text_lines(mar, lns, sz, lhh, "#555555")
        if it is not None:
            st["cv"].boxes.append((it, [mar, st["y"], mar + max(M.width(l, sz) for l in lns), st["y"] + h], s))
        transcript.append(s)
        st["y"] += h + 4

    new_page()
    if tpl.style == "table":
        sheets = _table_model(spec, tpl, items, tpl.roles)
        colw = [int((W - 2 * mar) * 0.38), int((W - 2 * mar) * 0.62)]
        for si, sh in enumerate(sheets):
            if si:
                ensure(lh + 14)
                st["y"] += 14
                text_lines(mar, [sh.name], size, lh, ink, True)
                transcript.append(sh.name)
                st["y"] += lh
            for r in sh.rows:
                if r.kind == "title":
                    ensure(lh + 6)
                    text_lines(mar, [r.cells[0]], size + 2, lh, ink, True)
                    transcript.append(r.cells[0])
                    st["y"] += lh + 6
                elif r.kind == "meta":
                    note(r.cells[0], r.item, size, lh)
                elif r.kind == "note":
                    note(r.cells[0], r.item, small, sl)
                else:
                    bold = r.kind == "hdr"
                    wr = [M.wrap(c, size, colw[j] - 20, bold) for j, c in enumerate(r.cells)]
                    h = max(len(w) for w in wr) * lh + 12
                    ensure(h)
                    cv, y, xx = st["cv"], st["y"], mar
                    cv.rect(mar, y, sum(colw), h, fill if bold else "#FFFFFF", "#9AA5B1")
                    for j, w in enumerate(wr):
                        if j:
                            cv.line(xx, y, xx, y + h, "#9AA5B1")
                        for k, ln in enumerate(w):
                            cv.text(xx + 10, y + 6 + k * lh, ln, size, ink, bold)
                        if r.item is not None and j == len(wr) - 1:
                            if r.item.keys == ("rate_tiers",):
                                # a tier is read as weight band plus price: the box covers the whole row
                                cv.boxes.append((r.item, [mar, y, mar + sum(colw), y + h], " | ".join(r.cells)))
                            else:
                                cv.boxes.append((r.item, [xx, y, xx + colw[j], y + h], r.cells[j]))
                        xx += colw[j]
                    transcript.append(" | ".join(r.cells))
                    st["y"] += h
        st["y"] += 10
    else:
        maxw = (W - 2 * mar) * (0.80 if tpl.style == "bubble" else 1.0) - (40 if tpl.style == "bubble" else 0)

        def block(s, it=None, small_=False, bold=False):
            sz = small if small_ else size
            lhh = sl if small_ else lh
            lns = M.wrap(s, sz, maxw, bold)
            h = len(lns) * lhh
            wmax = max(M.width(l, sz, bold) for l in lns)
            if tpl.style == "bubble":
                bw, bh = wmax + 2 * pad, h + 2 * pad
                ensure(bh + 10)
                bx = mar if tpl.align == "left" else W - mar - bw
                st["cv"].rect(bx, st["y"], bw, bh, fill, "#C4CCD6", 14)
                tx, ty, step = bx + pad, st["y"] + pad, bh + 10
            else:
                ensure(h + 6)
                tx, ty, step = mar + 16, st["y"], h + 6
            cv = st["cv"]
            for k, l in enumerate(lns):
                cv.text(tx, ty + k * lhh, l, sz, "#555555" if small_ else ink, bold)
            if it is not None:
                cv.boxes.append((it, [tx, ty, tx + wmax, ty + h], s))
            transcript.append(s)
            st["y"] += step

        block(ent_text, ent_item)
        for it in g["body"]:
            block(it.sentence, it)
        if g["annex"]:
            block(_m(spec, tpl.annex_label), None, False, True)
            for it in g["annex"]:
                block(it.sentence, it)
        for it in g["footnote"]:
            block(_m(spec, tpl.fn_label) + it.sentence, it, True)
        close = _closing(spec, tpl)
        if close:
            block(close, None, True)
    end_page()
    return pages, transcript


def _draw(cv: _Canvas, bg: str, path: Path, M: _Measure) -> None:
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure
    from matplotlib.patches import FancyBboxPatch, Rectangle

    W, H = cv.W, cv.H
    fig = Figure(figsize=(W / 100, H / 100), dpi=100, facecolor=bg)
    FigureCanvasAgg(fig)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, W)
    ax.set_ylim(H, 0)
    ax.axis("off")
    for op in cv.ops:
        if op[0] == "rect":
            _, x, y, w, h, fc, ec, r = op
            if r:
                ax.add_patch(FancyBboxPatch((x + r, y + r), w - 2 * r, h - 2 * r, boxstyle=f"round,pad={r}",
                                            fc=fc, ec=ec or fc, lw=1.2))
            else:
                ax.add_patch(Rectangle((x, y), w, h, fc=fc, ec=ec or fc, lw=1.0))
        elif op[0] == "line":
            ax.plot([op[1], op[3]], [op[2], op[4]], color=op[5], lw=1.0)
        else:
            _, x, y, s, size, color, bold = op
            ax.text(x, y, s, fontproperties=M.prop(size, bold), color=color, va="top", ha="left")
    fig.savefig(path, dpi=100, metadata={"Software": None})


def _render_image(spec: RenderSpec, tpl: Tpl, items: list, outdir: Path) -> tuple:
    base = _safe(spec.cand_id)
    M = _measure()
    pages = transcript = None
    size = tpl.size
    for size in range(tpl.size, MIN_FONT_PX - 1, -2):
        pages, transcript = _layout(spec, tpl, items, size, M, None)
        if len(pages) == 1 and max(pages[0].W, pages[0].H) <= MAX_EDGE:
            break
    else:
        size = MIN_FONT_PX
        pages, transcript = _layout(spec, tpl, items, size, M, MAX_EDGE)
    files, span_map = [], {}
    for i, cv in enumerate(pages):
        name = base + (".png" if i == 0 else f"_{i + 1}.png")
        _draw(cv, tpl.colors[0], outdir / name, M)
        assert max(cv.W, cv.H) <= MAX_EDGE, (name, cv.W, cv.H)
        files.append(str(outdir / name))
        for it, bb, quote in cv.boxes:
            box = [max(0, int(bb[0])), max(0, int(bb[1])), min(cv.W, int(bb[2]) + 1), min(cv.H, int(bb[3]) + 1)]
            _add_span(span_map, it, Span(name, "region", quote, bbox=box))
    layout = {"size": size, "W": tpl.width, "colors": tpl.colors, "align": tpl.align, "style": tpl.style,
              "font": M.info}
    meta = {"font_px": size, "note_px": max(size - 4, MIN_NOTE_PX), "pages": len(pages)}
    return files, "\n".join(transcript), span_map, {"pages": [cv.texts for cv in pages], "layout": layout}, meta


# ====================================================================== entry
def render(spec: RenderSpec, outdir: Any) -> RenderResult:
    """Render one candidate. image text_view is the transcript of the drawn text (for scoring only)."""
    spec.validate()
    tpl = TEMPLATES[spec.template_id]
    out = Path(outdir)
    out.mkdir(parents=True, exist_ok=True)
    items = _build_items(spec, tpl)
    fn = {"text": _render_text, "table": _render_table, "image": _render_image}[spec.modality]
    files, text, span_map, payload, meta = fn(spec, tpl, items, out)
    blob = json.dumps({"v": RENDER_VERSION, "spec": spec.to_dict(), "payload": payload},
                      sort_keys=True, ensure_ascii=False, default=str)
    return RenderResult(files, text, span_map, hashlib.sha256(blob.encode("utf-8")).hexdigest(), meta)
