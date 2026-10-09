"""Calculator B: landed cost computed by Excel formulas, recalculated with LibreOffice.

Python only writes inputs and formulas and reads back the recalculated cells; this module is
independent of calc_py by design (it imports schema and params only).

Conventions (same as calculator A):
- Every reported line is ROUNDed to 0.01 in the workbook; total = sum of the rounded lines;
  per_unit = total / sum(qty), rounded. Intermediate values are not rounded.
- Rate tier: the tier with the largest min_kg <= chargeable weight; max_kg is display only.
  A weight below the lowest min_kg raises ValueError.
- fuel_pct applies to freight only. pct_freight surcharges are fractions of freight.
- ctn_* surcharge conditions are only valid with basis per_ctn.
- DDP: "included" waives duty; scope duty_mpf and all also waive MPF; only scope all waives HMF.
  "capped" limits the seller's duty only (MPF and HMF stay with the seller).
- duty_floor_cny = duty + MPF (HMF excluded), before any DDP waiver.
- Missing optional terms are treated as zero/absent and listed in `defaulted`.
- Workbook files are named <shipment_id>_<input hash>.xlsx; each soffice call uses its own profile.
"""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
import uuid
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .params import Snapshot
from .schema import QuoteTerms, Shipment

SOFFICE_TIMEOUT_S = 180

# calculation sheet rows (row 1 is the header); K!B2 must stay chargeable_kg
_ROWS = [
    ("chargeable_kg", "=MAX(SUM(C!$G$2:$G${nc}),{min_kg_shipment})"),
    ("tier_hits", "=SUM(T!$G$2:$G${nt})"),
    ("tier_lo", "=SUMPRODUCT(MAX(T!$G$2:$G${nt}*T!$A$2:$A${nt}))"),
    ("price_cny", "=SUMIF(T!$A$2:$A${nt},{tier_lo},T!$F$2:$F${nt})"),
    ("freight_raw", "={chargeable_kg}*{price_cny}"),
    ("fuel_raw", "={fuel_pct}*{freight_raw}"),
    ("customs_value_usd", "=SUM(G!$E$2:$E${ng})"),
    ("duty_usd", "=SUM(G!$F$2:$F${ng})"),
    ("formal_flag", "=IF({customs_value_usd}>{formal_thr},1,0)"),
    ("sea_flag", '=IF({mode}="sea",1,0)'),
    ("mpf_usd", "=IF({formal_flag}=1,MIN(MAX({customs_value_usd}*{mpf_rate},{mpf_min}),{mpf_max}),{mpf_informal})"),
    ("hmf_usd", "=IF({sea_flag}=1,{customs_value_usd}*{hmf_rate},0)"),
    ("duty_cny", "={duty_usd}*{fx}"),
    ("mpf_cny", "={mpf_usd}*{fx}"),
    ("hmf_cny", "={hmf_usd}*{fx}"),
    ("duty_seller_raw", '=IF(AND({is_ddp}=1,{duty_terms}="included"),0,'
                        'IF(AND({is_ddp}=1,{duty_terms}="capped"),MAX(0,{duty_cny}-{ddp_cap}*{chargeable_kg}),{duty_cny}))'),
    ("mpf_seller_raw", '=IF(AND({is_ddp}=1,{duty_terms}="included",OR({ddp_scope}="duty_mpf",{ddp_scope}="all")),'
                       '0,{mpf_cny})'),
    ("hmf_seller_raw", '=IF(AND({is_ddp}=1,{duty_terms}="included",{ddp_scope}="all"),0,{hmf_cny})'),
    ("freight", "=ROUND(ROUND({freight_raw},9),2)"),
    ("fuel", "=ROUND(ROUND({fuel_raw},9),2)"),
    ("surcharges", "={surch_sum}"),
    ("contingent", "={cont_sum}"),
    ("duty", "=ROUND(ROUND({duty_seller_raw},9),2)"),
    ("mpf", "=ROUND(ROUND({mpf_seller_raw},9),2)"),
    ("hmf", "=ROUND(ROUND({hmf_seller_raw},9),2)"),
    ("broker", "=ROUND(ROUND({broker_in},9),2)"),
    ("isf", "=ROUND(ROUND({isf_in},9),2)"),
    ("total_cny", "=ROUND({freight}+{fuel}+{surcharges}+{duty}+{mpf}+{hmf}+{broker}+{isf},2)"),
    ("duty_floor_cny", "=ROUND(ROUND({duty_cny}+{mpf_cny},9),2)"),
    ("per_unit_cny", "=ROUND(ROUND({total_cny}/SUM(G!$B$2:$B${ng}),9),2)"),
]
_ROW_OF = {name: i for i, (name, _) in enumerate(_ROWS, start=2)}
_NAMES = [n for n, _ in _ROWS]


def _num(x):
    """Decimal -> float for a spreadsheet cell (data transfer only)."""
    return float(x) if isinstance(x, Decimal) else x


def _build(shipment: Shipment, terms: QuoteTerms, snap: Snapshot) -> Workbook:
    shipment.validate()
    terms.validate()
    if terms.plain("mode") is not None and terms.plain("mode") != shipment.mode:
        raise ValueError(f"terms mode {terms.plain('mode')!r} != shipment mode {shipment.mode!r}")
    for nm in ("dim_divisor", "rounding_kg", "rate_tiers"):
        if terms.plain(nm) is None:
            raise ValueError(f"quote terms missing required field: {nm}")

    wb = Workbook()
    P = wb.active
    P.title = "P"
    ddp = terms.plain("ddp")
    params = [
        ("mode", shipment.mode),
        ("dim_divisor", _num(terms.plain("dim_divisor"))),
        ("rounding_kg", _num(terms.plain("rounding_kg"))),
        ("min_kg_carton", _num(terms.plain("min_kg_carton") or 0)),
        ("min_kg_shipment", _num(terms.plain("min_kg_shipment") or 0)),
        ("fx", _num(snap.fx_cny_per_usd())),
        ("fuel_pct", _num(terms.plain("fuel_pct") or 0)),
        ("addr_type", shipment.dest.addr_type),
        ("remote", 1 if shipment.dest.remote else 0),
        ("mpf_rate", _num(snap.fee("mpf_rate"))),
        ("mpf_min", _num(snap.fee("mpf_min_usd"))),
        ("mpf_max", _num(snap.fee("mpf_max_usd"))),
        ("mpf_informal", _num(snap.fee("mpf_informal_usd"))),
        ("formal_thr", _num(snap.fee("formal_entry_threshold_usd"))),
        ("hmf_rate", _num(snap.fee("hmf_rate"))),
        ("is_ddp", 1 if (ddp is not None and ddp.is_ddp) else 0),
        ("ddp_scope", ddp.scope if ddp is not None else "duty"),
        ("duty_terms", (ddp.duty_terms or "actual") if ddp is not None else "actual"),
        ("ddp_cap", _num((ddp.cap_cny_per_kg if ddp is not None else None) or 0)),
        ("broker_in", _num(terms.plain("broker_fee_cny") or 0)),
        ("isf_in", _num(terms.plain("isf_fee_cny") or 0)),
    ]
    ref = {}
    for i, (k, v) in enumerate(params, start=1):
        P.cell(i, 1, k)
        P.cell(i, 2, v)
        ref[k] = f"P!$B${i}"
    K_CHARGE, K_FREIGHT = f"K!$B${_ROW_OF['chargeable_kg']}", f"K!$B${_ROW_OF['freight_raw']}"

    # cartons: n, L, W, H, kg | volume kg, per-carton chargeable kg (x n), longest side
    C = wb.create_sheet("C")
    C.append(["n", "L", "W", "H", "kg", "vol_kg", "charge_kg", "longest"])
    for i, c in enumerate(shipment.cartons, start=2):
        C.append([c.n, _num(c.length_cm), _num(c.width_cm), _num(c.height_cm), _num(c.kg)])
        C.cell(i, 6, f"=B{i}*C{i}*D{i}/{ref['dim_divisor']}")
        C.cell(i, 7, f"=A{i}*CEILING(ROUND(MAX(E{i},F{i},{ref['min_kg_carton']}),9),{ref['rounding_kg']})")
        C.cell(i, 8, f"=MAX(B{i},C{i},D{i})")
    nc = len(shipment.cartons) + 1
    nrng, kgrng, lrng = f"C!$A$2:$A${nc}", f"C!$E$2:$E${nc}", f"C!$H$2:$H${nc}"

    # rate tiers (the max column is carried for display only)
    T = wb.create_sheet("T")
    T.append(["min", "max", "price", "cur", "unused", "price_cny", "hit"])
    for i, t in enumerate(terms.plain("rate_tiers"), start=2):
        T.append([_num(t.min_kg), _num(t.max_kg), _num(t.price), t.currency])
        T.cell(i, 6, f'=C{i}*IF(D{i}="USD",{ref["fx"]},1)')
        T.cell(i, 7, f"=IF(A{i}<={K_CHARGE},1,0)")
    nt = len(terms.plain("rate_tiers")) + 1

    # duty component rates
    R = wb.create_sheet("R")
    R.append(["hts", "component", "rate"])
    seen = set()
    for g in shipment.goods:
        if g.hts in seen:
            continue
        seen.add(g.hts)
        for comp, rate, _cite, _st in snap.duty_components(g.hts):
            R.append([g.hts, comp, _num(rate)])
    nr = R.max_row

    # goods
    G = wb.create_sheet("G")
    G.append(["hts", "qty", "unit_usd", "rate", "value_usd", "duty_usd"])
    for i, g in enumerate(shipment.goods, start=2):
        G.append([g.hts, g.qty, _num(g.unit_value_usd)])
        G.cell(i, 4, f"=SUMIF(R!$A$2:$A${nr},A{i},R!$C$2:$C${nr})")
        G.cell(i, 5, f"=B{i}*C{i}")
        G.cell(i, 6, f"=E{i}*D{i}")
    ng = len(shipment.goods) + 1

    # surcharges: K raw value, L rounded, M counted in total, N contingent
    S = wb.create_sheet("S")
    S.append(["code", "basis", "amount", "cur", "cond", "thr", "contingent",
              "amt_cny", "ctn_count", "flag", "value", "rounded", "counted", "contingent_amt"])
    surcharges = terms.plain("surcharges") or []
    for i, s in enumerate(surcharges, start=2):
        S.append([s.code, s.basis, _num(s.amount), s.currency, s.condition, _num(s.threshold),
                  1 if s.contingent else 0])
        S.cell(i, 8, f'=C{i}*IF(D{i}="USD",{ref["fx"]},1)')
        S.cell(i, 9, f'=IF(E{i}="ctn_longest_side_gt",SUMPRODUCT({nrng}*({lrng}>F{i})),'
                     f'IF(E{i}="ctn_kg_gt",SUMPRODUCT({nrng}*({kgrng}>F{i})),SUM({nrng})))')
        S.cell(i, 10, f'=IF(E{i}="residential",IF({ref["addr_type"]}="residential",1,0),'
                      f'IF(E{i}="remote",{ref["remote"]},1))')
        S.cell(i, 11, f'=J{i}*IF(B{i}="per_kg",H{i}*{K_CHARGE},IF(B{i}="per_shpt",H{i},'
                      f'IF(B{i}="per_ctn",H{i}*I{i},C{i}*{K_FREIGHT})))')
        S.cell(i, 12, f"=ROUND(ROUND(K{i},9),2)")
        S.cell(i, 13, f"=L{i}*(1-G{i})")
        S.cell(i, 14, f"=L{i}*G{i}")
    ns = len(surcharges) + 1
    refs = dict(ref)
    refs["surch_sum"] = f"SUM(S!$M$2:$M${ns})" if surcharges else "0"
    refs["cont_sum"] = f"SUM(S!$N$2:$N${ns})" if surcharges else "0"
    for n in _NAMES:
        refs.setdefault(n, f"$B${_ROW_OF[n]}")
    # P-sheet references are absolute with sheet names; K-sheet names resolve in-sheet
    K = wb.create_sheet("K")
    K.append(["name", "value"])
    for name, tmpl in _ROWS:
        f = tmpl.replace("{nc}", str(nc)).replace("{nt}", str(nt)).replace("{ng}", str(ng))
        f = re.sub(r"\{(\w+)\}", lambda m: refs[m.group(1)], f)
        K.append([name, f])
    return wb


def _input_hash(shipment, terms, snap) -> str:
    blob = json.dumps({"s": shipment.to_dict(), "t": terms.to_dict(), "id": snap.snapshot_id,
                       "fees": snap._fees, "fx": snap._fx, "hts": snap.raw["hts"]}, sort_keys=True)
    return hashlib.sha256(blob.encode()).hexdigest()[:12]


def _recalc(src: Path, outdir: Path) -> Path:
    profile = outdir / f".lo_profile_{uuid.uuid4().hex}"
    cmd = ["soffice", "--headless", f"-env:UserInstallation={profile.resolve().as_uri()}",
           "--convert-to", "xlsx", "--outdir", str(outdir), str(src)]
    try:
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=SOFFICE_TIMEOUT_S)
    except (OSError, subprocess.TimeoutExpired) as e:
        raise RuntimeError(f"soffice failed: {e}") from e
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    out = outdir / src.name
    if r.returncode != 0 or not out.exists():
        raise RuntimeError(f"soffice recalculation failed rc={r.returncode}: {r.stderr[-300:]}")
    return out


def landed_cost_xlsx(shipment: Shipment, terms: QuoteTerms, snapshot: Snapshot, workdir) -> dict:
    """Returns chargeable_kg, total_cny, duty_floor_cny, per_unit_cny, lines [(code, amt)],
    contingent [(code, amt)], snapshot_id, assumptions, overrides, defaulted, fx_date, workbook."""
    workdir = Path(workdir)
    inputs = workdir / "inputs"
    inputs.mkdir(parents=True, exist_ok=True)
    wb = _build(shipment, terms, snapshot)
    name = re.sub(r"[^A-Za-z0-9_.-]", "_", shipment.shipment_id) + "_" + _input_hash(shipment, terms, snapshot)
    src = inputs / f"{name}.xlsx"
    wb.save(src)
    out = _recalc(src, workdir)
    wbr = load_workbook(out, data_only=True)
    ws = wbr["K"]
    vals = {ws.cell(i, 1).value: ws.cell(i, 2).value for i in range(2, ws.max_row + 1)}
    if any(vals.get(k) is None for k in _NAMES):
        raise RuntimeError("LibreOffice did not recalculate the workbook (empty result cells)")
    if vals["tier_hits"] < 1:
        raise ValueError("no rate tier covers the chargeable weight")
    d = {k: Decimal(repr(float(vals[k]))) for k in _NAMES}
    lines = [("freight", d["freight"]), ("fuel", d["fuel"])]
    contingent = []
    S = wbr["S"]
    for r in range(2, S.max_row + 1):
        code, amt, cont = S.cell(r, 1).value, S.cell(r, 12).value, S.cell(r, 7).value
        if amt is None:
            raise RuntimeError("LibreOffice did not recalculate the surcharge sheet")
        (contingent if cont == 1 else lines).append((code, Decimal(repr(float(amt)))))
    lines += [(k, d[k]) for k in ("duty", "mpf", "hmf", "broker", "isf")]
    formal, sea = vals["formal_flag"] == 1, vals["sea_flag"] == 1
    used_fees = ["formal_entry_threshold_usd"] + (["mpf_rate", "mpf_min_usd", "mpf_max_usd"] if formal
                                                  else ["mpf_informal_usd"]) + (["hmf_rate"] if sea else [])
    return {
        "chargeable_kg": d["chargeable_kg"],
        "total_cny": d["total_cny"],
        "duty_floor_cny": d["duty_floor_cny"],
        "per_unit_cny": d["per_unit_cny"],
        "lines": lines,
        "contingent": contingent,
        "snapshot_id": snapshot.snapshot_id,
        "assumptions": snapshot.assumed_among(used_fees, [g.hts for g in shipment.goods]),
        "overrides": snapshot.overrides(),
        "defaulted": terms.defaulted(),
        "fx_date": snapshot.fx_date() or "",
        "workbook": str(out),
    }
