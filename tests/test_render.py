import copy
import re
from pathlib import Path
from decimal import Decimal

import pytest

from quoteguard import render as R
from quoteguard.schema import Claim, DDP, Field, QuoteTerms, RateTier, Surcharge

ALL_TEMPLATES = sorted(R.TEMPLATES)


def terms_air_plain():
    return QuoteTerms.from_plain(
        channel="空派专线甲", mode="air",
        rate_tiers=[RateTier(21, 44, "42", "CNY"), RateTier(45, None, "39.5", "CNY")],
        min_kg_shipment=Decimal("21"), min_kg_carton=Decimal("1"), dim_divisor=Decimal("6000"),
        rounding_kg=Decimal("0.5"),
        surcharges=[Surcharge("REM", "per_kg", "3.2", "CNY", "remote"),
                    Surcharge("RES", "per_ctn", "25", "CNY", "residential", contingent=True),
                    Surcharge("OVS", "per_ctn", "60", "CNY", "ctn_longest_side_gt", "120")],
        fuel_pct=Decimal("0.12"), includes=["头程运费", "目的仓派送"], excludes=["关税", "MPF", "仓储费"],
        ddp=DDP(False, "duty", None), broker_fee_cny=Decimal("150"), transit_days=[6, 9],
        validity=["2026-10-12", "2026-10-31"], liability="按申报价值赔付，最高3倍运费")


def terms_air_claims():
    t = terms_air_plain()
    t.ddp = Field(DDP(True, "all", "included", None, "forwarder"), "stated", [])
    t.includes = Field(["头程运费", "关税", "MPF"], "stated", [])
    t.excludes = Field(["仓储费"], "stated", [])
    t.claims = Field([Claim("FMC_OTI", "SYN-Harbor Lane Logistics", "SYN-FMC-0001", "US"),
                      Claim("AEO_CN", "SYN-Harbor Lane Logistics", None, "CN"),
                      Claim("MOT_NVOCC", "SYN-Eastlake Logistics", None, "CN"),
                      Claim("NVOCC", "SYN-Bluewater Cargo", None, None)], "stated", [])
    return t


def terms_sea_conflict():
    t = QuoteTerms.from_plain(
        channel="海派快线乙", mode="sea",
        rate_tiers=[RateTier(0, 999, "1.85", "USD"), RateTier(1000, None, "1.6", "USD")],
        surcharges=[Surcharge("REM", "per_kg", "7", "CNY", "remote"), Surcharge("INSP", "per_shpt", "900", "CNY", "always", contingent=True)],
        fuel_pct=Decimal("0.08"), ddp=DDP(True, "duty_mpf", "capped", Decimal("4.5")),
        includes=["关税", "MPF"], validity=["2026-10-12", "2026-10-31"])
    t.transit_days = Field(None, "conflicting", [])
    t.surcharges = Field(None, "conflicting", [])
    return t


def terms_ddp_conflict():
    t = terms_air_plain()
    t.ddp = Field(None, "conflicting", [])
    t.includes = Field(["头程运费"], "stated", [])
    t.excludes = Field(["仓储费"], "stated", [])
    return t


SUR_A = [{"code": "REM", "basis": "per_kg", "amount": "7", "currency": "CNY", "condition": "remote"},
         {"code": "INSP", "basis": "per_shpt", "amount": "900", "currency": "CNY", "condition": "always",
          "contingent": True}]
SUR_B = [dict(SUR_A[0], amount="11"), SUR_A[1]]
CONFLICT_SEA = {"transit_days": [[28, 32], [35, 40]], "surcharges": [SUR_A, SUR_B]}
CONFLICT_DDP = {"ddp": [{"is_ddp": False, "scope": "duty", "duty_terms": "actual", "ior": "seller"},
                        {"is_ddp": True, "scope": "all", "duty_terms": "included", "ior": "forwarder"}]}
FEE_SEA = {"fuel_pct": "footnote", "surcharges": "annex"}
FEE_AIR = {"surcharges": "annex", "broker_fee_cny": "footnote"}

CASES = {
    "air_plain": (terms_air_plain, {}, FEE_AIR),
    "air_claims": (terms_air_claims, {}, {"surcharges": "annex"}),
    "sea_conflict": (terms_sea_conflict, CONFLICT_SEA, FEE_SEA),
    "ddp_conflict": (terms_ddp_conflict, CONFLICT_DDP, {"ddp": "footnote"}),
}


def per_view(spans):
    """spans per view: a table has a markdown (text) view and an xlsx (cell) view"""
    kinds = {}
    for sp in spans:
        kinds[sp.kind] = kinds.get(sp.kind, 0) + 1
    assert len(set(kinds.values())) == 1
    return next(iter(kinds.values()))


def make_spec(tid, case, **kw):
    mk, conf, fee = CASES[case]
    pool, mod = tid.split("-")[0], tid.split("-")[1]
    d = dict(cand_id=f"C-{case}", terms=mk(), modality=mod, template_id=tid, pool=pool,
             conflicts=copy.deepcopy(conf), fee_location=dict(fee), lang_mix=False,
             ddp_wording_id=f"{pool}-ddp-1", claim_wording_id=f"{pool}-claim-1",
             forwarder_alias="货代甲", quote_entity="SYN-Test Entity")
    d.update(kw)
    return R.RenderSpec(**d)


def expected_present(spec):
    out = {"quote_entity"}
    for n in R.CANON:
        f = getattr(spec.terms, n)
        if f.status != "missing" and f.value is not None:
            out.add(n)
    out |= set(spec.conflicts)
    return out


def test_template_inventory():
    for pool in R.POOLS:
        for mod in R.MODALITIES:
            ids = R.TEMPLATE_IDS[pool][mod]
            assert len(ids) >= 3
            assert ids == [f"{pool}-{mod}-{i}" for i in (1, 2, 3)]
        assert len(R.DDP_WORDING_IDS[pool]) >= 4 and len(R.CLAIM_WORDING_IDS[pool]) >= 4


def test_pools_share_no_sentence_pattern():
    dev, test = R.sentence_patterns("dev"), R.sentence_patterns("test")
    assert len(dev) > 50 and len(test) > 50
    assert not (dev & test), sorted(dev & test)


def test_field_orders_differ_between_pools():
    dev = {R.TEMPLATES[t].order for t in ALL_TEMPLATES if t.startswith("dev")}
    test = {R.TEMPLATES[t].order for t in ALL_TEMPLATES if t.startswith("test")}
    assert not (dev & test)


@pytest.mark.parametrize("case", list(CASES))
@pytest.mark.parametrize("tid", ALL_TEMPLATES)
def test_render_all(tid, case, tmp_path):
    spec = make_spec(tid, case)
    res = R.render(spec, tmp_path)
    assert res.files and all((tmp_path / f.split("/")[-1]).exists() for f in res.files)
    assert res.render_hash and res.text_view
    keys = {k for k in res.span_map if "." not in k}
    assert expected_present(spec) <= keys
    for n in spec.conflicts:
        if n != "surcharges":
            assert per_view(res.span_map[n]) == 2
    if "surcharges" in spec.conflicts:
        assert per_view(res.span_map["surcharges.REM"]) == 2      # value differs: two spans
        assert per_view(res.span_map["surcharges.INSP"]) == 1     # unchanged: once
    assert per_view(res.span_map["quote_entity"]) == 1
    assert "SYN-Test Entity" in res.text_view
    mod = spec.modality
    for key, spans in res.span_map.items():
        for sp in spans:
            sp.validate()
            if mod == "image":
                assert sp.kind == "region"
                assert sp.quote in res.text_view
            elif sp.kind == "text":
                assert res.text_view[sp.start:sp.end] == sp.quote
    if mod == "table":
        from openpyxl import load_workbook
        wb = load_workbook([f for f in res.files if f.endswith(".xlsx")][0])
        cells = [sp for sps in res.span_map.values() for sp in sps if sp.kind == "cell"]
        assert cells
        for sp in cells:
            assert wb[sp.sheet][sp.cell].value == sp.quote
        # every text span has a matching cell span count
        n_text = sum(1 for sps in res.span_map.values() for sp in sps if sp.kind == "text")
        assert n_text == len(cells)
    if mod == "image":
        from PIL import Image
        im = Image.open([f for f in res.files if f.endswith(".png")][0])
        W, H = im.size
        assert max(W, H) <= 1600
        for sps in res.span_map.values():
            for sp in sps:
                x0, y0, x1, y1 = sp.bbox
                assert 0 <= x0 < x1 <= W and 0 <= y0 < y1 <= H
    if mod == "text":
        assert all(sp.kind == "text" for sps in res.span_map.values() for sp in sps)


@pytest.mark.parametrize("tid", ALL_TEMPLATES)
def test_missing_value_absent_and_present_when_stated(tid, tmp_path):
    spec = make_spec(tid, "air_plain")
    spec.terms.isf_fee_cny = Field(Decimal("7.91"), "stated", [])
    spec.terms.broker_fee_cny = Field(Decimal("150"), "stated", [])
    spec.terms.liability = Field(None, "missing", [])
    r1 = R.render(spec, tmp_path / "a")
    assert "7.91" in r1.text_view and "isf_fee_cny" in r1.span_map
    spec.terms.isf_fee_cny = Field(None, "missing", [])
    r2 = R.render(spec, tmp_path / "b")
    assert "7.91" not in r2.text_view and "isf_fee_cny" not in r2.span_map
    assert "liability" not in r2.span_map and "赔付" not in r2.text_view


@pytest.mark.parametrize("tid", ALL_TEMPLATES)
def test_conflicting_values_both_present(tid, tmp_path):
    spec = make_spec(tid, "sea_conflict")
    res = R.render(spec, tmp_path)
    tv = res.text_view
    for tok in ("28", "32", "35", "40", "7", "11", "900"):
        assert tok in tv
    q = [sp.quote for sp in res.span_map["transit_days"] if sp.kind != "cell"]
    assert len(set(q)) == 2
    q2 = [sp.quote for sp in res.span_map["surcharges.REM"] if sp.kind != "cell"]
    assert len(set(q2)) == 2
    # the unchanged line is written once
    assert tv.count("900") == 1
    # missing fields stay out
    assert "isf_fee_cny" not in res.span_map and "dim_divisor" not in res.span_map


@pytest.mark.parametrize("tid", ["dev-text-1", "dev-table-2", "dev-image-3", "test-text-3", "test-table-1",
                                 "test-image-2"])
def test_hash_stable_and_sensitive(tid, tmp_path):
    a = R.render(make_spec(tid, "air_claims"), tmp_path / "a")
    b = R.render(make_spec(tid, "air_claims"), tmp_path / "b")
    assert a.render_hash == b.render_hash and a.text_view == b.text_view
    assert a.span_map.keys() == b.span_map.keys()
    c = R.render(make_spec(tid, "air_claims", lang_mix=True), tmp_path / "c")
    assert c.render_hash != a.render_hash


def test_lang_mix_adds_english(tmp_path):
    for tid in ("dev-text-1", "test-text-1"):
        res = R.render(make_spec(tid, "air_plain", lang_mix=True), tmp_path)
        tv = res.text_view
        assert "per" in tv or "/KG" in tv
        assert "Remote area" in tv or "Outlying zone" in tv
        assert "  " not in tv and not re.search(r"[A-Za-z]{3,}[\u4e00-\u9fff]|[\u4e00-\u9fff][A-Za-z]{3,}", tv), tv
        assert all(l == l.rstrip() for l in tv.split("\n"))
        plain = R.render(make_spec(tid, "air_plain"), tmp_path)
        assert "Remote" not in plain.text_view


def test_files_are_ordered_by_template_and_fee_location(tmp_path):
    res = R.render(make_spec("dev-text-1", "air_plain"), tmp_path)
    tv = res.text_view
    i_note, i_annex = tv.index("备注："), tv.index("附加费说明：")
    resi = res.span_map["surcharges.RES"][0]
    over = res.span_map["surcharges.OVS"][0]
    assert resi.start > i_annex and over.start > i_annex
    assert res.span_map["broker_fee_cny"][0].start > i_note
    assert res.span_map["rate_tiers"][0].start < i_note
    assert len(res.span_map["rate_tiers"]) == 2 and len(res.span_map["surcharges"]) == 3


@pytest.mark.parametrize("pool", R.POOLS)
def test_ddp_wording_matches_terms(pool, tmp_path):
    lx = R.LEX[pool][False]
    for wid in R.DDP_WORDING_IDS[pool]:
        w = lx["ddp_w"][wid]
        tid = f"{pool}-text-1"
        sp = make_spec(tid, "air_plain", ddp_wording_id=wid)
        sp.terms.includes = Field(None, "missing", [])
        sp.terms.excludes = Field(None, "missing", [])
        assert w["none"] in R.render(sp, tmp_path).text_view
        for terms_ddp, cls in ((DDP(True, "all", "included"), "included"),
                               (DDP(True, "duty", "actual"), "actual"),
                               (DDP(True, "duty_mpf", "capped", Decimal("4.5")), "capped")):
            sp.terms.ddp = Field(terms_ddp, "stated", [])
            tv = R.render(sp, tmp_path).text_view
            assert w[cls].format(cap="").split("{")[0] in tv or w[cls].split("{cap}")[0] in tv
            assert w["none"] not in tv
        if pool == "dev" and wid == "dev-ddp-3":
            assert "DDP 包税" in R.render(make_spec(tid, "air_claims", ddp_wording_id=wid), tmp_path).text_view
    for wid in R.CLAIM_WORDING_IDS[pool]:
        res = R.render(make_spec(f"{pool}-text-2", "air_claims", claim_wording_id=wid), tmp_path)
        assert "SYN-Harbor Lane Logistics" in res.text_view and "SYN-FMC-0001" in res.text_view
        assert len(res.span_map["claims"]) == 4
        assert not re.search(r"_CN|/CN|_OTI|MOT_|AEO_|FMC_", res.text_view)


def test_unknown_ids_rejected(tmp_path):
    with pytest.raises(ValueError):
        R.render(make_spec("dev-text-1", "air_plain", template_id="dev-text-9"), tmp_path)
    with pytest.raises(ValueError):
        R.render(make_spec("dev-text-1", "air_plain", ddp_wording_id="test-ddp-1"), tmp_path)
    with pytest.raises(ValueError):
        R.render(make_spec("dev-text-1", "air_plain", template_id="test-text-1"), tmp_path)
    spec = make_spec("dev-text-1", "sea_conflict", conflicts={})
    with pytest.raises(ValueError):
        R.render(spec, tmp_path)


def test_spec_from_dict_roundtrip(tmp_path):
    pres = {"modality": "text", "template_id": "test-text-2", "pool": "test", "lang_mix": True,
            "ddp_wording_id": "test-ddp-2", "claim_wording_id": "test-claim-3",
            "fee_location": {"fuel_pct": "footnote"}, "conflicts": {}, "forwarder_alias": "货代乙",
            "quote_entity": "SYN-Zed Cargo"}
    terms = terms_air_claims()
    spec = R.RenderSpec.from_dict(pres, cand_id="K1", terms=terms.to_dict())
    assert isinstance(spec.terms, QuoteTerms) and spec.terms.to_dict() == terms.to_dict()
    res = R.render(spec, tmp_path)
    assert "货代乙" in res.text_view and res.span_map["fuel_pct"][0].start > res.text_view.index("注：")
    with pytest.raises(ValueError):
        R.RenderSpec.from_dict(pres)


@pytest.mark.parametrize("tid", ALL_TEMPLATES)
@pytest.mark.parametrize("ior,word", [("forwarder", 0), ("seller", 1)])
def test_ior_rendered_with_span(tid, ior, word, tmp_path):
    spec = make_spec(tid, "air_claims")
    d = spec.terms.ddp.value
    spec.terms.ddp = Field(DDP(d.is_ddp, d.scope, d.duty_terms, None, ior), "stated", [])
    res = R.render(spec, tmp_path)
    w = R.LEX[spec.pool][False]["ior_w"][ior]
    sps = res.span_map["ddp.ior"]
    assert sps and all(w in sp.quote for sp in sps)
    assert "ddp" in res.span_map
    spec.terms.ddp = Field(DDP(d.is_ddp, d.scope, d.duty_terms, None, "unspecified"), "stated", [])
    assert "ddp.ior" not in R.render(spec, tmp_path).span_map


def test_ior_wordings_differ_between_pools():
    assert not (set(R.LEX["dev"][False]["ior_w"].values()) & set(R.LEX["test"][False]["ior_w"].values()))
    assert R.LEX["dev"][True]["S"]["ior"] != R.LEX["dev"][False]["S"]["ior"]


def test_wrap_keeps_ascii_runs_whole():
    M = R._Measure()
    s = "前缀前缀前缀前缀FBA仓含关税和MPF 以及 ONT8 仓 per kg 计费2026-10-12后缀文字"
    ascii_ = lambda c: c.isascii() and c.isalnum()
    for maxw in range(150, 520, 13):
        lines = M.wrap(s, 24, maxw)
        assert "".join(lines) == s
        for a, b in zip(lines, lines[1:]):
            assert not (a and b and ascii_(a[-1]) and ascii_(b[0])), (maxw, a, b)
    assert max(M.width(l, 24) for l in M.wrap("A" * 80, 24, 200)) <= 200 + 30


# ---------------------------------------------------------------- taskgen-driven checks
@pytest.fixture(scope="module")
def dev_tasks():
    from quoteguard import taskgen
    tasks, _, _ = taskgen.generate_split("dev", seed=7, n_per_layer=3)
    return tasks


def _spec_from_candidate(cand, **over):
    """Spec from a taskgen candidate; tax words in includes/excludes are aligned with ddp.scope
    (a no-op once the generator follows the rule)."""
    spec = R.RenderSpec.from_dict(cand.presentation, **over)
    t = spec.terms
    d = t.ddp.value if t.ddp.status == "stated" else None
    universe = [tx for tx in R.TAXES if tx != "HMF" or t.mode.value == "sea"]
    want = (R._SCOPE_TAXES[d.scope] & set(universe)) if d is not None and d.is_ddp else set()
    for nm in ("includes", "excludes"):
        f = getattr(t, nm)
        if f.status != "stated" or not f.value:
            continue
        keep = [x for x in f.value if not any(tx in x for tx in R.TAXES)]
        if nm == "includes":
            keep += [tx for tx in universe if tx in want]
        elif d is not None and not d.is_ddp:
            keep += universe
        setattr(t, nm, Field(keep, "stated", []))
    return spec


def _retarget(cand, pool, modality, idx=1, wid=1):
    return {"pool": pool, "modality": modality, "template_id": f"{pool}-{modality}-{idx}",
            "ddp_wording_id": f"{pool}-ddp-{wid}", "claim_wording_id": f"{pool}-claim-{wid}"}


def _strip_data(text, pres):
    """Remove everything that comes from the quote data rather than the wording."""
    terms = pres["terms"]
    data = [pres["quote_entity"], pres["forwarder_alias"]]
    for n in ("channel", "liability"):
        v = terms[n]["value"]
        if v:
            data.append(v)
    for n in ("includes", "excludes"):
        data += list(terms[n]["value"] or [])
    for c in terms["claims"]["value"] or []:
        data += [c["entity"], c.get("reg_no") or ""]
    data += list(R.TAXES)
    for d in sorted((x for x in data if x), key=len, reverse=True):
        text = text.replace(d, " ")
    return text


def _grams(text, n=4):
    out = set()
    for line in text.split("\n"):
        norm = re.sub(r"[^0-9A-Za-z一-鿿]", "", line.lower())
        norm = re.sub(r"[0-9]", "", norm)
        out |= {norm[i:i + n] for i in range(len(norm) - n + 1)}
    return out


def test_pool_texts_share_no_4gram(dev_tasks, tmp_path):
    corpus = {"dev": set(), "test": set()}
    k = 0
    for t in dev_tasks:
        for c in t.candidates:
            for pool in R.POOLS:
                for mix in (False, True):
                    for mod, idx in (("text", 1 + k % 3), ("table", 1 + (k + 1) % 3)):
                        k += 1
                        over = _retarget(c, pool, mod, idx, 1 + k % 4)
                        over["lang_mix"] = mix
                        if mod == "table":
                            over["fee_location"] = c.presentation["fee_location"]
                        try:
                            res = R.render(_spec_from_candidate(c, **over), tmp_path)
                        except ValueError as e:
                            raise AssertionError(f"{t.task_id}/{c.cand_id}: {e}")
                        corpus[pool] |= _grams(_strip_data(res.text_view, c.presentation))
    both = corpus["dev"] & corpus["test"]
    assert len(corpus["dev"]) > 200 and len(corpus["test"]) > 200
    assert not both, sorted(both)


def test_vocabulary_covers_generator_codes_and_claim_types():
    import inspect
    from pathlib import Path
    import yaml
    from quoteguard import taskgen
    data = Path(taskgen.__file__).resolve().parents[2] / "data" / "taskgen"
    structures = yaml.safe_load((data / "structures.yaml").read_text(encoding="utf-8"))
    value_sets = yaml.safe_load((data / "value_sets.yaml").read_text(encoding="utf-8"))
    codes = set(value_sets.get("surcharge_amounts", {}))

    def walk(o):
        if isinstance(o, dict):
            if "code" in o:
                codes.add(o["code"])
            for v in o.values():
                walk(v)
        elif isinstance(o, list):
            for v in o:
                walk(v)

    walk(structures)
    types = set(structures["claims"]["fmc_types"])
    toks = set(re.findall(r'"([A-Z][A-Z_]+)"', inspect.getsource(taskgen._draw_claim)))
    types |= {x for x in toks if len(x) > 2}
    juris = {x for x in toks if len(x) == 2}
    assert {"RES", "REM", "OVS", "OVW", "INSP"} <= codes
    assert {"FMC_OTI", "MOT_NVOCC", "AEO_CN", "NVOCC"} <= types
    for pool in R.POOLS:
        lx = R.LEX[pool][False]
        assert codes <= set(lx["names"]), codes - set(lx["names"])
        assert types <= set(lx["types"]), types - set(lx["types"])
        assert juris and juris <= set(lx["juris"])
        for v in list(lx["types"].values()) + list(lx["juris"].values()):
            assert not re.search(r"_CN|/CN|[A-Z]+_[A-Z]+", v)
    # no fallback to the raw code
    sp = make_spec("dev-text-1", "air_plain")
    sp.terms.surcharges = Field([Surcharge("ZZZ", "per_kg", "1", "CNY", "always")], "stated", [])
    with pytest.raises(ValueError):
        R.render(sp, "/tmp/_qg_unused")
    sp = make_spec("dev-text-1", "air_claims")
    sp.terms.claims = Field([Claim("XYZ_CN", "SYN-A", None, None)], "stated", [])
    with pytest.raises(ValueError):
        R.render(sp, "/tmp/_qg_unused")


@pytest.mark.parametrize("tid", ALL_TEMPLATES)
def test_quote_entity_written_in_every_template(tid, tmp_path):
    res = R.render(make_spec(tid, "air_plain", quote_entity="SYN-Quoter One"), tmp_path)
    assert "SYN-Quoter One" in res.text_view
    (sp,) = [s for s in res.span_map["quote_entity"] if s.kind != "cell"]
    assert "SYN-Quoter One" in sp.quote
    with pytest.raises(ValueError):
        R.render(make_spec(tid, "air_plain", quote_entity=""), tmp_path)


def test_tax_items_follow_ddp_scope(tmp_path):
    def spec(ddp, inc, exc, mode="sea"):
        sp = make_spec("dev-text-1", "air_plain")
        sp.terms.mode = Field(mode, "stated", [])
        sp.terms.ddp = ddp
        sp.terms.includes = Field(inc, "stated", [])
        sp.terms.excludes = Field(exc, "stated", [])
        return sp

    d = lambda scope, ddp=True: Field(DDP(ddp, scope, "included" if ddp else None), "stated", [])
    ok = [spec(d("duty"), ["头程运费", "关税"], ["MPF", "HMF"]),
          spec(d("duty_mpf"), ["关税", "MPF"], ["HMF"]),
          spec(d("all"), ["关税", "MPF", "HMF"], ["仓储费"]),
          spec(d("duty", False), ["头程运费"], ["关税", "MPF", "HMF"])]
    for sp in ok:
        R.render(sp, tmp_path)
    bad = [spec(d("duty"), ["关税", "MPF"], []),                     # includes more than scope
           spec(d("all"), ["关税"], []),                             # includes less than scope
           spec(d("duty", False), ["关税"], ["MPF"]),                 # not DDP but taxes included
           spec(d("duty", False), [], ["关税"]),                      # not DDP: all three must be excluded
           spec(d("duty"), ["关税"], ["关税"])]                       # tax both included and excluded
    for sp in bad:
        with pytest.raises(ValueError):
            R.render(sp, tmp_path)
    # ddp missing or conflicting: no tax word anywhere
    sp = spec(Field(None, "missing", []), ["头程运费", "关税"], [])
    with pytest.raises(ValueError):
        R.render(sp, tmp_path)
    sp = make_spec("dev-text-1", "ddp_conflict")
    sp.terms.excludes = Field(["仓储费", "MPF"], "stated", [])
    with pytest.raises(ValueError):
        R.render(sp, tmp_path)
    sp = spec(Field(None, "missing", []), ["头程运费"], ["仓储费"])
    R.render(sp, tmp_path)


@pytest.mark.parametrize("mode", ["air", "express"])
def test_hmf_only_in_sea_quotes(mode, tmp_path):
    def spec(ddp, inc, exc):
        sp = make_spec("dev-text-1", "air_plain")
        sp.terms.mode = Field(mode, "stated", [])
        sp.terms.ddp = Field(ddp, "stated", [])
        sp.terms.includes = Field(inc, "stated", [])
        sp.terms.excludes = Field(exc, "stated", [])
        return sp

    R.render(spec(DDP(True, "all", "included"), ["关税", "MPF"], ["仓储费"]), tmp_path)      # all = duty and MPF
    R.render(spec(DDP(True, "duty", "included"), ["关税"], ["MPF"]), tmp_path)
    R.render(spec(DDP(False, "duty", None), [], ["关税", "MPF"]), tmp_path)                  # no HMF to exclude
    for sp in (spec(DDP(True, "all", "included"), ["关税", "MPF", "HMF"], []),
               spec(DDP(True, "duty", "included"), ["关税"], ["MPF", "HMF"]),
               spec(DDP(False, "duty", None), [], ["关税", "MPF", "HMF"]),
               spec(DDP(True, "all", "included"), ["关税"], [])):
        with pytest.raises(ValueError):
            R.render(sp, tmp_path)
    sp = spec(DDP(True, "all", "included"), ["关税", "MPF"], ["仓储费"])
    sp.terms.mode = Field("sea", "stated", [])          # the same lists are one tax short for a sea quote
    with pytest.raises(ValueError):
        R.render(sp, tmp_path)


def test_conflict_validation(tmp_path):
    sp = make_spec("dev-text-1", "air_plain", conflicts={"fuel_pct": [0.1, 0.2]})
    with pytest.raises(ValueError):           # field is not conflicting
        R.render(sp, tmp_path)
    sp = make_spec("dev-text-1", "sea_conflict")
    sp.terms.fuel_pct = Field(None, "conflicting", [])
    with pytest.raises(ValueError):           # conflicting without an entry
        R.render(sp, tmp_path)
    sp = make_spec("dev-text-1", "sea_conflict", conflicts={**CONFLICT_SEA, "transit_days": [[28, 32], [28, 32]]})
    with pytest.raises(ValueError):           # two equal values
        R.render(sp, tmp_path)
    sp = make_spec("dev-text-1", "sea_conflict", fee_location={"nonexistent": "annex"})
    with pytest.raises(ValueError):
        R.render(sp, tmp_path)
    sp = make_spec("dev-text-1", "sea_conflict", fee_location={"surcharges.REM": "annex"})
    with pytest.raises(ValueError):
        R.render(sp, tmp_path)
    # the candidates stored in terms must equal conflicts
    sp = make_spec("dev-text-1", "sea_conflict")
    sp.terms.transit_days = Field([[28, 32], [36, 40]], "conflicting", [])
    sp._conf_raw["transit_days"] = [[28, 32], [36, 40]]
    with pytest.raises(ValueError):
        R.render(sp, tmp_path)
    sp = make_spec("dev-text-1", "sea_conflict")
    sp._conf_raw["transit_days"] = [[28, 32], [35, 40]]
    R.render(sp, tmp_path)
    # a presentation dict with conflicting values inside terms
    pres = {"modality": "text", "template_id": "dev-text-1", "pool": "dev", "quote_entity": "SYN-X",
            "conflicts": copy.deepcopy(CONFLICT_SEA), "fee_location": dict(FEE_SEA)}
    td = terms_sea_conflict().to_dict()
    td["transit_days"]["value"] = [[28, 32], [35, 40]]
    td["surcharges"]["value"] = [SUR_A, SUR_B]
    res = R.render(R.RenderSpec.from_dict(pres, cand_id="D1", terms=td), tmp_path)
    assert len(res.span_map["transit_days"]) == 2
    td["transit_days"]["value"] = [[28, 32], [35, 41]]
    with pytest.raises(ValueError):
        R.render(R.RenderSpec.from_dict(pres, cand_id="D1", terms=td), tmp_path)


def test_surcharge_conflict_writes_changed_line_twice(tmp_path):
    res = R.render(make_spec("dev-text-1", "sea_conflict"), tmp_path)
    tv = res.text_view
    assert len(res.span_map["surcharges.REM"]) == 2 and len(res.span_map["surcharges.INSP"]) == 1
    assert len(res.span_map["surcharges"]) == 3
    assert tv.count("查验费") == 1 and tv.count("偏远地区附加费") == 2


def test_zero_values(tmp_path):
    def spec(tid, mode, **zero):
        sp = make_spec(tid, "air_plain")
        sp.terms.mode = Field(mode, "stated", [])
        if mode == "sea":
            sp.terms.excludes = Field(["关税", "MPF", "HMF", "仓储费"], "stated", [])
        sp.terms.fuel_pct = Field(Decimal(0), "stated", [])
        sp.terms.isf_fee_cny = Field(Decimal(0), "stated", [])
        sp.terms.broker_fee_cny = Field(Decimal(0), "stated", [])
        sp.terms.min_kg_carton = Field(Decimal(0), "stated", [])
        sp.terms.min_kg_shipment = Field(Decimal(0), "stated", [])
        sp.terms.surcharges = Field([Surcharge("REM", "per_kg", "0", "CNY", "remote")], "stated", [])
        return sp

    for pool in R.POOLS:
        tid = f"{pool}-text-1"
        lx = R.LEX[pool][False]
        r = R.render(spec(tid, "express"), tmp_path)
        assert lx["S"]["fuel_pct0"] in r.text_view and "fuel_pct" in r.span_map
        assert lx["S"]["min_kg_shipment0"] in r.text_view
        for k in ("isf_fee_cny", "broker_fee_cny", "min_kg_carton", "surcharges"):
            assert k not in r.span_map, k
        r = R.render(spec(tid, "sea"), tmp_path)
        assert lx["S"]["isf_fee_cny0"] in r.text_view
        assert "fuel_pct" not in r.span_map
        r = R.render(spec(tid, "air"), tmp_path)
        assert "fuel_pct" not in r.span_map and "isf_fee_cny" not in r.span_map


def test_ior_follows_ddp_location(tmp_path):
    for tid in ("dev-text-1", "test-table-2", "dev-image-2", "test-image-2"):
        sp = make_spec(tid, "air_claims", fee_location={"ddp": "footnote"})
        sp.terms.ddp = Field(DDP(True, "all", "included", None, "seller"), "stated", [])
        res = R.render(sp, tmp_path)
        a = [s for s in res.span_map["ddp"] if s.kind != "cell"][0]
        b = [s for s in res.span_map["ddp.ior"] if s.kind != "cell"][0]
        if sp.modality == "image":
            tv = res.text_view
            assert tv.index(a.quote) > tv.index(res.span_map["rate_tiers"][0].quote)
            assert tv.index(b.quote) > tv.index(res.span_map["rate_tiers"][0].quote)
            assert abs(tv.index(a.quote) - tv.index(b.quote)) < 120      # adjacent notes
        else:
            anchor = res.text_view.index("备注：" if sp.pool == "dev" else "注：")
            assert a.start >= anchor and b.start >= anchor


def test_table_header_rows_not_repeated(tmp_path):
    from openpyxl import load_workbook
    for tid in R.TEMPLATE_IDS["dev"]["table"] + R.TEMPLATE_IDS["test"]["table"]:
        res = R.render(make_spec(tid, "sea_conflict"), tmp_path)
        wb = load_workbook([f for f in res.files if f.endswith(".xlsx")][0])
        for ws in wb:
            rows = [tuple(c.value for c in r) for r in ws.iter_rows()]
            assert len(rows) == len(set(rows)), (tid, ws.title)


def test_tier_rows_boxed_whole_in_table_images(tmp_path):
    res = R.render(make_spec("dev-image-2", "air_plain"), tmp_path)
    tiers = res.span_map["rate_tiers"]
    assert len(tiers) == 2
    for sp in tiers:
        assert sp.bbox[2] - sp.bbox[0] > 1000 and " | " in sp.quote and "kg" in sp.quote and "CNY" not in sp.quote
        assert sp.quote in res.text_view
    other = res.span_map["fuel_pct"][0]
    assert other.bbox[2] - other.bbox[0] < 900


def test_card_has_inner_padding(tmp_path):
    res = R.render(make_spec("dev-image-3", "air_plain"), tmp_path)
    first = min(sp.bbox[1] for sps in res.span_map.values() for sp in sps)
    assert first >= 64 + 14 + 18      # header bar, card border, padding


# ---------------------------------------------------------------- byte level determinism
@pytest.mark.parametrize("tid", ["dev-text-2", "dev-table-1", "test-table-3", "dev-image-1", "test-image-3"])
def test_files_are_byte_identical(tid, tmp_path):
    a = R.render(make_spec(tid, "sea_conflict"), tmp_path / "a")
    import time
    time.sleep(1.1)
    b = R.render(make_spec(tid, "sea_conflict"), tmp_path / "b")
    assert len(a.files) == len(b.files)
    for fa, fb in zip(a.files, b.files):
        assert Path(fa).read_bytes() == Path(fb).read_bytes(), fa


# ---------------------------------------------------------------- images: font floor and pages
@pytest.mark.parametrize("pool", R.POOLS)
def test_all_dev_candidates_render_on_every_image_template(pool, dev_tasks, tmp_path):
    sizes, pages = [], 0
    for t in dev_tasks:
        for c in t.candidates:
            for tid in R.TEMPLATE_IDS[pool]["image"]:
                over = _retarget(c, pool, "image", int(tid[-1]), 1)
                spec = _spec_from_candidate(c, **over)
                res = R.render(spec, tmp_path)
                assert res.meta["font_px"] >= R.MIN_FONT_PX and res.meta["note_px"] >= R.MIN_NOTE_PX
                sizes.append(res.meta["font_px"])
                pages += res.meta["pages"] - 1
                from PIL import Image
                names = {Path(f).name for f in res.files}
                for f in res.files:
                    assert max(Image.open(f).size) <= 1600
                for sps in res.span_map.values():
                    for sp in sps:
                        assert sp.doc_id in names and sp.kind == "region"
    assert min(sizes) >= R.MIN_FONT_PX


def test_long_quote_splits_into_second_image(tmp_path):
    sp = make_spec("dev-image-1", "air_claims")
    sp.terms.surcharges = Field([Surcharge(c, "per_kg", "3.2", "CNY", "always")
                                 for c in ("RES", "REM", "OVS", "OVW", "INSP")], "stated", [])
    sp.fee_location = {"surcharges": "annex"}
    sp.terms.includes = Field(["头程运费", "关税", "MPF"], "stated", [])
    sp.terms.liability = Field("按申报价值赔付，" * 60, "stated", [])
    res = R.render(sp, tmp_path)
    assert res.meta["pages"] >= 2 and res.meta["font_px"] == R.MIN_FONT_PX
    assert Path(res.files[1]).name == "C-air_claims_2.png"
    docs = {sp_.doc_id for sps in res.span_map.values() for sp_ in sps}
    assert docs == {Path(f).name for f in res.files}


# ---------------------------------------------------------------- font
def test_missing_font_raises_and_font_enters_hash(tmp_path, monkeypatch):
    spec = make_spec("dev-image-1", "air_plain")
    h1 = R.render(spec, tmp_path).render_hash
    monkeypatch.setattr(R, "_M", None)
    monkeypatch.setattr(R, "_FONT_FAMILY", "No Such CJK Font")
    with pytest.raises(RuntimeError):
        R.render(spec, tmp_path)
    monkeypatch.setattr(R, "_FONT_FAMILY", "Noto Sans CJK SC")
    monkeypatch.setattr(R, "_M", None)
    assert R.render(spec, tmp_path).render_hash == h1
    real = R._font_info
    monkeypatch.setattr(R, "_M", None)
    monkeypatch.setattr(R, "_font_info", lambda: {**real(), "matplotlib": "0.0"})
    assert R.render(spec, tmp_path).render_hash != h1
    monkeypatch.setattr(R, "_M", None)


@pytest.mark.parametrize("tid", ALL_TEMPLATES)
def test_zero_on_one_side_of_a_conflict_is_written(tid, tmp_path):
    pool = tid.split("-")[0]
    lx = R.LEX[pool][False]
    for field_, vals, key in (("broker_fee_cny", ["0", "150"], "broker_fee_cny0"),
                              ("fuel_pct", ["0", "0.1"], "fuel_pct0"),
                              ("isf_fee_cny", ["200", "0"], "isf_fee_cny0"),
                              ("min_kg_carton", ["0", "2"], "min_kg_carton0")):
        sp = make_spec(tid, "air_plain", fee_location={})
        setattr(sp.terms, field_, Field(None, "conflicting", []))
        sp.conflicts = {field_: vals}
        res = R.render(sp, tmp_path)
        assert len(res.span_map[field_]) >= 2 and per_view(res.span_map[field_]) == 2
        zero_pat = lx["S"][key]
        lab, val = lx["L"][key]
        assert zero_pat in res.text_view or (lab in res.text_view and val in res.text_view)
        quotes = [s.quote for s in res.span_map[field_]]
        assert any(zero_pat in q or q == val for q in quotes), quotes
    # surcharge table with a zero amount on one side
    sp = make_spec(tid, "sea_conflict")
    zero_b = [dict(SUR_A[0], amount="0"), SUR_A[1]]
    sp.conflicts = {**CONFLICT_SEA, "surcharges": [SUR_A, zero_b]}
    res = R.render(sp, tmp_path)
    assert per_view(res.span_map["surcharges.REM"]) == 2
    zq = [s.quote for s in res.span_map["surcharges.REM"]]
    assert any(lx["S"]["surcharge0"].format(name=lx["names"]["REM"]) in q or q == lx["L"]["surcharge0"][1]
               for q in zq), zq
    sp.conflicts = {**CONFLICT_SEA, "surcharges": [zero_b, SUR_A]}
    res = R.render(sp, tmp_path)
    assert per_view(res.span_map["surcharges.REM"]) == 2


@pytest.mark.parametrize("tid", [t for t in ALL_TEMPLATES if "-image-" in t])
def test_long_tier_list_paginates_within_limit(tid, tmp_path):
    from PIL import Image
    sp = make_spec(tid, "air_plain")
    sp.terms.rate_tiers = Field([RateTier(i * 10, i * 10 + 9, str(30 + i % 7), "CNY") for i in range(40)],
                                "stated", [])
    res = R.render(sp, tmp_path)
    assert res.meta["pages"] >= 2
    for f in res.files:
        assert max(Image.open(f).size) <= 1600
    assert len(res.span_map["rate_tiers"]) == 40


def test_mixed_text_has_no_space_before_fullwidth_punctuation(tmp_path):
    tv = R.render(make_spec("dev-text-1", "air_plain", lang_mix=True), tmp_path).text_view
    assert not re.search(r" [，。：；）]", tv)
