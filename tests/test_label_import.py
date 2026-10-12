"""标签导入工具的测试：全部使用合成的单元、匿名件、初标与已填裁决表/复核表，不含任何真实裁决。"""
import hashlib
import json
import shutil
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import import_adjudication as adj  # noqa: E402
import import_review13 as r13  # noqa: E402
import label_common as lc  # noqa: E402

from quoteguard.schema import LabelRecord  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DOC_A = "合成渠道 SYN-空派甲\n单箱不低于 12KG，按 12KG 计费。\n材积 长*宽*高/6000\n报关费 RMB300/票\n备注 报关费 350 元/票\n价格 美西 30 元/KG"
DOC_B = "合成渠道 SYN-海派乙  海派 价格 8 元/KG 起\n含税\n时效 20-25 天"


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def meta(sid, src, **kw):
    d = {"seed_id": sid, "source_id": src, "source_type": "web_page", "source_url": f"{src}-src-1",
         "captured_at": "2026-10-09", "page_date": None, "forwarder_alias": "F90", "publisher_alias": "F90",
         "modality": "text", "channel": "SYN", "raw_sha256": "0" * 64, "anon_sha256": None, "is_real": True,
         "source_class": "forwarder", "price_form": "published_tier", "date_status": "unknown",
         "dest_region": "美西"}
    d.update(kw)
    return d


def F(status, value=None, *quotes, doc="S01"):
    return {"status": status, "value": value, "spans": [{"doc_id": doc, "kind": "text", "quote": q} for q in quotes]}


@pytest.fixture
def env(tmp_path):
    anon = tmp_path / "anon"
    anon.mkdir()
    (anon / "S01.txt").write_text(DOC_A, encoding="utf-8")
    (anon / "S02.txt").write_text(DOC_B, encoding="utf-8")
    units = [meta("S01-u1", "S01", anon_sha256=sha(anon / "S01.txt")),
             meta("S01-u2", "S01", anon_sha256=sha(anon / "S01.txt")),
             meta("S02-u1", "S02", anon_sha256=sha(anon / "S02.txt"))]
    (tmp_path / "units.json").write_text(json.dumps(units, ensure_ascii=False), encoding="utf-8")
    (tmp_path / "blind.json").write_text(json.dumps({"selected_ids": ["S01-u1", "S01-u2"]}), encoding="utf-8")
    return tmp_path


def initial(sid, k):
    """两份合成初标：min_kg_carton、broker_fee_cny 与 page_date 不同，其余相同。"""
    terms = {
        "channel": F("stated", "SYN-空派甲", "合成渠道 SYN-空派甲"),
        "mode": F("implied", "air", "SYN-空派甲"),
        "min_kg_carton": F("stated", "12" if k == 1 else "6", "单箱不低于 12KG"),
        "dim_divisor": F("stated", "6000", "材积 长*宽*高/6000"),
        "broker_fee_cny": (F("stated", "300", "报关费 RMB300/票") if k == 1 else
                           F("conflicting", ["300", "350"], "报关费 RMB300/票", "备注 报关费 350 元/票")),
        "rate_tiers": F("stated", [{"min_kg": "0", "max_kg": None, "price": "30.0" if k == 1 else "30",
                                    "currency": "CNY", "unit": "per_kg"}], "价格 美西 30 元/KG"),
    }
    for n in lc.FIELDS:
        terms.setdefault(n, F("missing"))
    return {"seed_id": sid, "annotator": f"init{k}", "annotated_at": "2026-10-11", "minutes_spent": "0",
            "meta": {"page_date": None if k == 1 else "2026-10-01", "price_form": "published_tier",
                     "date_status": "unknown"},
            "terms": terms, "notes": {"src": "合成", "only1": "x"} if k == 1 else {"src": "合成"}}


def write_initials(env):
    for k in (1, 2):
        (env / f"init{k}.json").write_text(json.dumps([initial(s, k) for s in ("S01-u1", "S01-u2")],
                                                      ensure_ascii=False), encoding="utf-8")


HEAD = ["seed_id", "字段", "初标1", "初标2", "决定", "裁决 status", "裁决 value", "原文依据", "理由", "裁决人", "日期"]


def sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    ws.append([f"{title}（合成测试）"])
    ws.append(HEAD)
    for r in rows:
        ws.append([r.get(h) for h in HEAD])


def diff_rows(sid, **over):
    rows = {"min_kg_carton": {"决定": "采用初标1"},
            "broker_fee_cny": {"决定": "采用初标2"}}
    rows.update(over)
    return [{"seed_id": sid, "字段": n, "裁决人": "裁决人甲", "日期": "2026-10-12", **v} for n, v in rows.items()]


def agree_rows(sid, skip=()):
    return [{"seed_id": sid, "字段": n, "决定": "确认", "裁决人": "裁决人甲", "日期": "2026-10-12"}
            for n in lc.FIELDS if n not in ("min_kg_carton", "broker_fee_cny") and n not in skip]


def make_xlsx(env, diff=None, agree=None, metas=None, notes=None):
    wb = Workbook()
    wb.active.title = "说明"
    sheet(wb, "差异裁决", diff if diff is not None else diff_rows("S01-u1") + diff_rows("S01-u2"))
    sheet(wb, "一致项复核", agree if agree is not None else agree_rows("S01-u1") + agree_rows("S01-u2"))
    sheet(wb, "元数据差异", metas if metas is not None else [
        {"seed_id": s, "字段": "page_date", "决定": "改正", "裁决 value": "2026-10-02", "理由": "合成",
         "裁决人": "裁决人甲", "日期": "2026-10-12"} for s in ("S01-u1", "S01-u2")])
    sheet(wb, "单元备注", notes if notes is not None else [
        {"seed_id": s, "字段": "only1", "决定": "删除", "理由": "合成"} for s in ("S01-u1", "S01-u2")])
    p = env / "裁决表.xlsx"
    wb.save(p)
    return p


def adj_args(env, *extra):
    return ["--xlsx", str(env / "裁决表.xlsx"), "--initial-1", str(env / "init1.json"),
            "--initial-2", str(env / "init2.json"), "--units", str(env / "units.json"),
            "--blind", str(env / "blind.json"), "--anon", str(env / "anon"),
            "--labels-dir", str(env / "labels"), "--adjudication", str(env / "adjudication.json"), *extra]


# ---------------------------------------------------------------- (a) 人工裁决表
def test_adjudication_happy_path_writes_gold_and_record(env):
    write_initials(env)
    diff = diff_rows("S01-u1") + diff_rows("S01-u2", min_kg_carton={
        "决定": "改正", "裁决 status": "stated", "裁决 value": "12", "原文依据": "[S01] 按 12KG 计费", "理由": "合成"})
    make_xlsx(env, diff=diff)
    assert adj.main(adj_args(env)) == 0
    g = json.loads((env / "labels/S01-u2.gold.json").read_text(encoding="utf-8"))
    LabelRecord.from_dict(g).validate()
    assert g["annotator"] == "gold" and g["annotated_at"] == "2026-10-12"
    assert g["terms"]["min_kg_carton"]["value"] == "12"
    assert g["terms"]["min_kg_carton"]["spans"][0]["quote"] == "按 12KG 计费"
    assert g["terms"]["broker_fee_cny"]["status"] == "conflicting"
    assert g["terms"]["broker_fee_cny"]["value"] == ["300", "350"]
    assert g["meta"]["page_date"] == "2026-10-02" and g["meta"]["anon_sha256"] == sha(env / "anon/S01.txt")
    assert g["notes"] == {"src": "合成"}
    a = json.loads((env / "adjudication.json").read_text(encoding="utf-8"))
    s = a["summary"]
    assert s["fields"] == 36 and s["initials_differ"] == 4
    assert s["decisions"] == {"adopted_1": 1, "adopted_2": 2, "confirmed": 32, "corrected": 1}
    # 初标1 在两个单元的 broker_fee_cny 被改；初标2 在两个单元的 min_kg_carton 被改
    assert s["correction_rate"]["initial_1"]["corrected"] == 2
    assert s["correction_rate"]["initial_2"]["corrected"] == 2
    assert a["gold_files"]["S01-u1.gold.json"] == sha(env / "labels/S01-u1.gold.json")
    assert len(a["notes"]) == 2


def test_omitted_default_keys_do_not_count_as_corrections(env):
    """初标档位省略 unit（schema 默认 per_kg）时，与金标签相同，不计入改正。"""
    write_initials(env)
    for k in (1, 2):
        p = env / f"init{k}.json"
        recs = json.loads(p.read_text(encoding="utf-8"))
        for r in recs:
            del r["terms"]["rate_tiers"]["value"][0]["unit"]
        p.write_text(json.dumps(recs, ensure_ascii=False), encoding="utf-8")
    make_xlsx(env)
    assert adj.main(adj_args(env)) == 0
    a = json.loads((env / "adjudication.json").read_text(encoding="utf-8"))
    assert a["summary"]["correction_rate"]["initial_1"]["corrected"] == 2


def test_empty_evidence_quote_is_rejected(env, capsys):
    write_initials(env)
    make_xlsx(env, diff=diff_rows("S01-u1", min_kg_carton={"决定": "改正", "裁决 status": "stated",
                                                           "裁决 value": "99", "原文依据": "[S01]"})
              + diff_rows("S01-u2"))
    assert adj.main(adj_args(env)) == 1
    assert "出处片段为空" in capsys.readouterr().err


def test_unrecognized_yellow_column_is_an_error(env, capsys):
    write_initials(env)
    make_xlsx(env)
    wb = load_workbook(env / "裁决表.xlsx")
    wb["差异裁决"].cell(2, 8).value = "原文依据（逐字）"
    wb.save(env / "裁决表.xlsx")
    assert adj.main(adj_args(env)) == 1
    assert "找不到列 原文依据" in capsys.readouterr().err


def test_meta_confirm_on_differing_initials_bad_key_and_note_problems(env, capsys):
    write_initials(env)
    metas = [{"seed_id": "S01-u1", "字段": "page_date", "决定": "确认", "理由": "合成", "裁决人": "甲", "日期": "2026-10-12"},
             {"seed_id": "S01-u2", "字段": "page_date", "决定": "采用初标1", "裁决人": "甲", "日期": "2026-10-12"},
             {"seed_id": "S01-u2", "字段": "price form", "决定": "改正", "裁决 value": "range", "裁决人": "甲",
              "日期": "2026-10-12"}]
    notes = [{"seed_id": "S01-u1", "字段": "only1", "决定": "删除"}, {"seed_id": "S01-u1", "字段": "only1", "决定": "保留"}]
    make_xlsx(env, metas=metas, notes=notes)
    assert adj.main(adj_args(env)) == 1
    err = capsys.readouterr().err
    assert "S01-u1.page_date" in err and "两份初标此项不同" in err
    assert "'price form' 不是 SeedQuote" in err
    assert "S01-u1.only1）: 重复" in err
    assert "S01-u2.notes.only1" in err


def test_numeric_formatting_differences_count_as_agreement(env):
    write_initials(env)          # rate_tiers 价格 30.0 与 30
    make_xlsx(env)
    assert adj.main(adj_args(env, "--check-only")) == 0


def test_missing_row_and_wrong_sheet_are_reported_and_nothing_written(env, capsys):
    write_initials(env)
    agree = agree_rows("S01-u1", skip=("validity",)) + agree_rows("S01-u2")
    diff = diff_rows("S01-u1") + diff_rows("S01-u2") + [
        {"seed_id": "S01-u1", "字段": "validity", "决定": "确认", "裁决人": "裁决人甲", "日期": "2026-10-12"}]
    make_xlsx(env, diff=diff, agree=agree)
    assert adj.main(adj_args(env)) == 1
    err = capsys.readouterr().err
    assert "S01-u1.validity" in err and "差异裁决" in err
    assert not (env / "labels").exists() and not (env / "adjudication.json").exists()


def test_unfilled_decision_and_adjudicator_are_reported(env, capsys):
    write_initials(env)
    diff = diff_rows("S01-u1", min_kg_carton={"决定": None}) + diff_rows("S01-u2")
    diff[-1]["裁决人"] = None
    make_xlsx(env, diff=diff)
    assert adj.main(adj_args(env)) == 1
    err = capsys.readouterr().err
    assert "决定未填" in err and "裁决人未填" in err


def test_adopt_with_different_value_is_rejected(env, capsys):
    write_initials(env)
    make_xlsx(env, diff=diff_rows("S01-u1", min_kg_carton={"决定": "采用初标1", "裁决 value": "10"})
              + diff_rows("S01-u2"))
    assert adj.main(adj_args(env)) == 1
    assert "改为“改正”" in capsys.readouterr().err


def test_correction_needs_evidence_and_evidence_must_be_substring(env, capsys):
    write_initials(env)
    make_xlsx(env, diff=diff_rows("S01-u1", min_kg_carton={"决定": "改正", "裁决 status": "stated",
                                                           "裁决 value": "12"})
              + diff_rows("S01-u2", min_kg_carton={"决定": "改正", "裁决 status": "stated", "裁决 value": "12",
                                                   "原文依据": "单箱不低于 15KG"}))
    assert adj.main(adj_args(env)) == 1
    err = capsys.readouterr().err
    assert "原文依据必填" in err
    # 第一个问题修好后，第二个单元的片段不在匿名件中
    make_xlsx(env, diff=diff_rows("S01-u1") + diff_rows("S01-u2", min_kg_carton={
        "决定": "改正", "裁决 status": "stated", "裁决 value": "12", "原文依据": "单箱不低于 15KG"}))
    assert adj.main(adj_args(env)) == 1
    assert "出处不是 S01 的子串" in capsys.readouterr().err


def test_schema_violation_in_corrected_value_is_reported(env, capsys):
    write_initials(env)
    make_xlsx(env, diff=diff_rows("S01-u1", broker_fee_cny={
        "决定": "改正", "裁决 status": "conflicting", "裁决 value": '["300"]', "原文依据": "报关费 RMB300/票"})
        + diff_rows("S01-u2"))
    assert adj.main(adj_args(env)) == 1
    assert "schema 校验失败" in capsys.readouterr().err


def test_initial_span_that_no_longer_matches_the_anonymized_file_fails(env, capsys):
    """匿名件补遮后，初标中引用了被遮内容的片段在出处核对中报错。"""
    write_initials(env)
    make_xlsx(env)
    p = env / "anon/S01.txt"
    p.write_text(DOC_A.replace("材积 长*宽*高/6000", "材积 [已匿名]"), encoding="utf-8")
    units = json.loads((env / "units.json").read_text(encoding="utf-8"))
    for u in units:
        if u["source_id"] == "S01":
            u["anon_sha256"] = sha(p)
    (env / "units.json").write_text(json.dumps(units, ensure_ascii=False), encoding="utf-8")
    assert adj.main(adj_args(env)) == 1
    assert "dim_divisor" in capsys.readouterr().err


def test_anon_hash_mismatch_is_reported(env, capsys):
    write_initials(env)
    make_xlsx(env)
    (env / "anon/S01.txt").write_text(DOC_A + "\n", encoding="utf-8")
    assert adj.main(adj_args(env)) == 1
    assert "anon_sha256 不一致" in capsys.readouterr().err


def test_initials_may_be_directories_and_must_cover_the_blind_subset(env, capsys):
    for k in (1, 2):
        d = env / f"init{k}"
        d.mkdir()
        (d / "S01-u1.json").write_text(json.dumps(initial("S01-u1", k), ensure_ascii=False), encoding="utf-8")
    make_xlsx(env)
    args = adj_args(env)
    args[args.index(str(env / "init1.json"))] = str(env / "init1")
    args[args.index(str(env / "init2.json"))] = str(env / "init2")
    assert adj.main(args) == 1
    assert "缺少单元：['S01-u2']" in capsys.readouterr().err


def test_inspect_lists_headers(env, capsys):
    make_xlsx(env)
    assert adj.main(["--xlsx", str(env / "裁决表.xlsx"), "--inspect"]) == 0
    out = capsys.readouterr().out
    assert "[差异裁决]" in out and "表头行：2" in out and "裁决 status" in out


# ---------------------------------------------------------------- (b) 13 单元复核表
def review_env(env):
    lab = env / "r13/labels"
    units_dir = env / "r13/units"
    lab.mkdir(parents=True)
    rec = initial("S02-u1", 1)
    rec["terms"] = {n: F("missing") for n in lc.FIELDS}
    rec["terms"]["channel"] = F("stated", "SYN-海派乙", "合成渠道 SYN-海派乙", doc="S02")
    rec["terms"]["mode"] = F("stated", "sea", "海派", doc="S02")
    rec["terms"]["transit_days"] = F("stated", [20, 26], "时效 20-25 天", doc="S02")
    rec.update({"seed_id": "S02-u1", "annotator": "claude_initial", "initial_label_model": "Claude（合成测试）",
                "meta": meta("S02-u1", "S02"), "notes": {"price_form": "starting_price"}})
    (lab / "S02-u1.claude_initial.json").write_text(json.dumps(rec, ensure_ascii=False), encoding="utf-8")
    sys.path.insert(0, str(ROOT / "scripts"))
    import make_review_pack as mrp
    d = units_dir / "S02-u1"
    d.mkdir(parents=True)
    wb = Workbook()
    mrp.sheet(wb, rec)
    wb.save(d / "S02-u1_复核表.xlsx")
    return d / "S02-u1_复核表.xlsx"


def fill(xlsx, overrides):
    wb = load_workbook(xlsx)
    ws = wb.active
    for r in range(4, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if name in lc.FIELDS:
            vals = overrides.get(name, ["确认"])
            for i, v in enumerate(vals):
                ws.cell(r, 6 + i).value = v
    wb.save(xlsx)


def r13_args(env, *extra):
    return ["--reviewer", "复核人甲", "--date", "2026-10-13", "--review-dir", str(env / "r13/units"),
            "--initial-dir", str(env / "r13/labels"), "--units", str(env / "units.json"),
            "--blind", str(env / "blind.json"), "--anon", str(env / "anon"),
            "--labels-dir", str(env / "labels"), "--record", str(env / "review13.json"), *extra]


def test_review13_happy_path(env):
    x = review_env(env)
    fill(x, {"transit_days": ["改正", "stated", "[20, 25]", "时效 20-25 天", "上限抄错"],
             "ddp": ["改正", "implied", '{"is_ddp": true, "scope": "duty", "duty_terms": "included", '
                     '"cap_cny_per_kg": null, "ior": "unspecified"}', "[S02] 含税", "表头写含税"]})
    assert r13.main(r13_args(env)) == 0
    g = json.loads((env / "labels/S02-u1.gold.json").read_text(encoding="utf-8"))
    LabelRecord.from_dict(g).validate()
    assert g["terms"]["transit_days"]["value"] == [20, 25]
    assert g["terms"]["ddp"]["value"]["is_ddp"] is True
    assert g["terms"]["channel"]["spans"][0]["quote"] == "合成渠道 SYN-海派乙"
    assert g["initial_label_model"] == "Claude（合成测试）" and g["annotated_at"] == "2026-10-13"
    rec = json.loads((env / "review13.json").read_text(encoding="utf-8"))
    assert rec["summary"]["corrected"] == 2 and rec["summary"]["value_or_status_changed"] == 2
    assert rec["summary"]["changed_fields"] == ["S02-u1.ddp", "S02-u1.transit_days"]
    assert rec["reviewer"] == "复核人甲"


def test_review13_confirm_with_filled_correction_is_rejected(env, capsys):
    x = review_env(env)
    fill(x, {"transit_days": ["确认", None, "[20, 25]"]})
    assert r13.main(r13_args(env)) == 1
    assert "选了“确认”却填了" in capsys.readouterr().err


def test_review13_unfilled_or_unreasoned_rows_fail(env, capsys):
    x = review_env(env)
    fill(x, {"mode": [None], "transit_days": ["改正", "stated", "[20, 25]", "时效 20-25 天", None]})
    assert r13.main(r13_args(env)) == 1
    err = capsys.readouterr().err
    assert "S02-u1.mode" in err and "理由必填" in err
    assert not (env / "review13.json").exists()


@pytest.mark.skipif(shutil.which("pdftotext") is None, reason="需要 pdftotext")
def test_real_review_pack_spans_match_current_anonymized_files(tmp_path):
    """仓库中 13 个单元的 Claude 初标出处与当前匿名件一致（全部按“确认”填写的临时副本，不写入仓库）。"""
    src = ROOT / "output/label_review_13/units"
    dst = tmp_path / "units"
    shutil.copytree(src, dst)
    for x in dst.glob("*/*_复核表.xlsx"):
        fill(x, {})
    rc = r13.main(["--reviewer", "检查", "--date", "2026-10-12", "--check-only", "--review-dir", str(dst)])
    assert rc == 0


# ---------------------------------------------------------------- 盲标一致性抽检
import import_spotcheck as spot  # noqa: E402

SPOT_HEAD = ["编号", "单元", "原件", "渠道/目的区域", "字段", "字段含义", "status", "value", "币种", "原文依据", "备注"]


def spot_xlsx(env, rows):
    wb = Workbook()
    wb.active.title = "说明"
    ws = wb.create_sheet("抽检")
    ws.append(["盲标字段抽检表（合成测试）"])
    ws.append(SPOT_HEAD)
    for r in rows:
        ws.append([r.get(h) for h in SPOT_HEAD])
    p = env / "抽检表.xlsx"
    wb.save(p)
    return p


def spot_args(env, *extra):
    return ["--xlsx", str(env / "抽检表.xlsx"), "--seed", "20261012", "--checker", "抽检人甲", "--date", "2026-10-13",
            "--blind", str(env / "blind.json"), "--labels-dir", str(env / "labels"), "--out", str(env / "spot.json"),
            *extra]


def test_spotcheck_rates_status_and_value_separately(env):
    write_initials(env)
    make_xlsx(env)
    assert adj.main(adj_args(env)) == 0          # 合成金标签：min_kg_carton=12，broker_fee_cny 冲突 [300, 350]
    spot_xlsx(env, [
        {"编号": "C01", "单元": "S01-u1", "字段": "min_kg_carton", "status": "stated", "value": "12.0"},
        {"编号": "C02", "单元": "S01-u1", "字段": "broker_fee_cny", "status": "stated", "value": "300"},
        {"编号": "C03", "单元": "S01-u2", "字段": "rate_tiers", "status": "stated", "币种": "CNY",
         "value": '[{"min_kg": "0", "max_kg": null, "price": "30"}]'},
        {"编号": "C04", "单元": "S01-u2", "字段": "fuel_pct", "status": "missing"},
        {"编号": "C05", "单元": "S01-u2", "字段": "dim_divisor", "status": "implied", "value": "5000"},
    ])
    assert spot.main(spot_args(env)) == 0
    r = json.loads((env / "spot.json").read_text(encoding="utf-8"))
    s = r["summary"]
    assert s["status"] == {"agree": 3, "n": 5, "rate": 0.6}       # C02、C05 status 不同
    assert s["value"] == {"agree": 3, "n": 5, "rate": 0.6}        # C02、C05 value 不同
    assert s["value_gold_present"] == {"agree": 2, "n": 3, "rate": 0.6667}
    assert s["disagreements"] == ["C02", "C05"]
    assert r["sampling"]["seed"] == 20261012 and r["inputs"]["gold_files"]


def test_spotcheck_input_errors(env, capsys):
    write_initials(env)
    make_xlsx(env)
    assert adj.main(adj_args(env)) == 0
    spot_xlsx(env, [
        {"编号": "C01", "单元": "S01-u1", "字段": "min_kg_carton", "status": "missing", "value": "12"},
        {"编号": "C01", "单元": "S01-u1", "字段": "mode", "status": "stated", "value": "air"},
        {"编号": "C03", "单元": "S09-u1", "字段": "mode", "status": "stated", "value": "air"},
        {"编号": "C04", "单元": "S01-u1", "字段": "单箱最低计费重", "status": "stated", "value": "12"},
        {"编号": "C05", "单元": "S01-u1", "字段": "mode", "status": None},
    ])
    assert spot.main(spot_args(env)) == 1
    err = capsys.readouterr().err
    for s in ("value 应留空", "编号重复", "不在盲标名单", "未知条款字段", "status 须为"):
        assert s in err
    assert not (env / "spot.json").exists()


def test_spotcheck_requires_gold_and_columns(env, capsys):
    spot_xlsx(env, [{"编号": "C01", "单元": "S01-u1", "字段": "mode", "status": "stated", "value": "air"}])
    assert spot.main(spot_args(env)) == 1
    assert "先运行 import_adjudication.py" in capsys.readouterr().err
    wb = load_workbook(env / "抽检表.xlsx")
    wb["抽检"].cell(2, 9).value = "货币"
    wb.save(env / "抽检表.xlsx")
    assert spot.main(spot_args(env)) == 1
    assert "找不到列 ['币种']" in capsys.readouterr().err
